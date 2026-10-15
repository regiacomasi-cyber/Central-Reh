"""
Analisador Curva ABC — simples de usar, resultado profissional.
Passo 1: mande sua planilha | Passo 2: veja sua análise | Passo 3: monte seus anúncios.
100% local, linguagem simples. Rodar: python app.py (ou pythonw app.py p/ sem janela preta)
"""
import csv
import json
import os
import warnings
import webbrowser
import urllib.request
import urllib.error
from urllib.parse import quote_plus


def http_get(url, timeout=20):
    req = urllib.request.Request(url, headers={
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        charset = r.headers.get_content_charset() or "utf-8"
        return r.read().decode(charset, errors="ignore")
warnings.filterwarnings("ignore")
try:
    import tkinter as tk
    from tkinter import ttk, scrolledtext, messagebox, filedialog, simpledialog
except ImportError:  # servidor web sem interface gráfica: só a lógica é usada
    import types
    tk = types.SimpleNamespace(Tk=object, Canvas=object)
    ttk = scrolledtext = messagebox = filedialog = simpledialog = None

# ================= identidade visual =================
# Identidade Central REH: bordô profundo + dourado + creme (estética premium)
# (nomes internos AZUL/LARANJA mantidos para não quebrar o código)
AZUL = "#3D0A1A"; LARANJA = "#C9A227"; FUNDO = "#FAF7F0"; CINZA_TXT = "#2B2320"
ROSA_CLARO = "#FBF1DC"; ROSA_ESCURO = "#2E0712"; OURO = "#E3B341"; CREME = "#FAF3E6"
JANELA_LARG, JANELA_ALT = 1120, 800

class BotaoRedondo(tk.Canvas):
    """Botão moderno com cantos arredondados e efeito ao passar o mouse."""
    def __init__(self, master, texto, comando, largura=230, altura=46, fundo=LARANJA,
                 cor_texto="#1A1A1A", fonte=("Segoe UI", 12, "bold")):
        try:
            bg = master.cget("background")
        except Exception:
            bg = FUNDO
        super().__init__(master, width=largura, height=altura, bg=bg,
                         highlightthickness=0, bd=0, cursor="hand2")
        self.fundo = fundo
        self._desenhar(fundo)
        self.create_text(largura//2, altura//2, text=texto, font=fonte, fill=cor_texto)
        self.bind("<Enter>", lambda e: self._desenhar(self._escurecer(fundo)))
        self.bind("<Leave>", lambda e: self._desenhar(fundo))
        self.bind("<Button-1>", lambda e: comando())

    def _desenhar(self, cor):
        self.delete("fundo")
        w, h = int(self.cget("width")), int(self.cget("height"))
        r = 16
        self.create_rectangle(r, 0, w-r, h, fill=cor, outline=cor, tags="fundo")
        self.create_rectangle(0, r, w, h-r, fill=cor, outline=cor, tags="fundo")
        for x, y in ((r, r), (w-r, r), (r, h-r), (w-r, h-r)):
            self.create_oval(x-r, y-r, x+r, y+r, fill=cor, outline=cor, tags="fundo")
        self.lower("fundo")

    @staticmethod
    def _escurecer(cor):
        try:
            v = [max(int(cor[i:i+2], 16)-24, 0) for i in (1, 3, 5)]
            return "#%02x%02x%02x" % tuple(v)
        except Exception:
            return cor

# ================= leitura de arquivos =================
COL_PROD = ["produto", "produtos", "nome", "sku", "item", "descricao", "descrição"]
COL_QTD = ["quantidade", "qtd", "qde", "vendidos", "vendas", "unidades", "pecas", "peças",
           "volume de vendas valido", "volume", "volume de vendas"]
COL_FAT = ["faturamento", "receita", "valor total", "total", "faturado", "venda total",
           "valor de vendas validas", "valor de vendas", "vendas validas"]
COL_PRECO = ["preco", "preço", "valor", "preco unit", "preço unit", "unitario", "unitário",
             "preco medio", "preço médio"]
COL_EST = ["estoque", "saldo", "estoque atual", "quantidade estoque"]
COL_CUSTO = ["custo", "cmv", "custo unit", "custo unitario", "custo unitário", "preco custo", "preço custo"]

def norm(s):
    return (s or "").strip().lower().replace("ã", "a").replace("õ", "o").replace("é", "e") \
        .replace("ê", "e").replace("ç", "c").replace("í", "i").replace("ó", "o").replace("á", "a")

def num(x):
    x = str(x or "").strip().replace("R$", "").replace(" ", "").replace("%", "")
    if not x or x == "auto": return 0.0
    return float(x.replace(".", "").replace(",", ".") if "," in x else x.replace(",", ""))

def achar(row, candidatos):
    for k, v in row.items():
        if norm(k) in candidatos: return v or ""
    return ""

def _montar_lista(linhas_dict):
    out = []
    for row in linhas_dict:
        prod = (achar(row, COL_PROD) or "").strip()
        if not prod: continue
        qtd = num(achar(row, COL_QTD))
        fat = num(achar(row, COL_FAT))
        if fat <= 0:
            fat = num(achar(row, COL_PRECO)) * qtd
        out.append({"produto": prod, "qtd": int(qtd), "fat": round(fat, 2),
                    "est": int(num(achar(row, COL_EST))), "custo": round(num(achar(row, COL_CUSTO)), 2)})
    if not out:
        raise ValueError("Não entendi sua planilha. Ela precisa ter colunas de produto, quantidade e faturamento (ou preço).")
    return out

def _ler_texto(caminho):
    """Lê CSV em qualquer codificação comum (acentos do Excel BR inclusos)."""
    for enc in ("utf-8-sig", "latin-1", "cp1252"):
        try:
            with open(caminho, encoding=enc) as f:
                return f.read()
        except (UnicodeDecodeError, ValueError):
            continue
    raise ValueError("Não consegui entender o texto do arquivo.")

def ler_csv(caminho):
    txt = _ler_texto(caminho)
    linhas = [l for l in txt.splitlines() if l.strip()]
    if not linhas: raise ValueError("Arquivo vazio.")
    primeira = linhas[0]
    delim = ";" if primeira.count(";") >= primeira.count(",") else ","
    rd = csv.DictReader(linhas, delimiter=delim)
    if not rd.fieldnames: raise ValueError("Arquivo vazio ou sem títulos na primeira linha.")
    return _montar_lista(list(rd))

def _norm_cab(cab):
    return [norm(c) for c in cab]

def eh_relatorio_pedidos(cab):
    n = _norm_cab(cab)
    return "id do pedido" in n and "nome do produto" in n

def _idx(cab, nomes):
    n = _norm_cab(cab)
    for alvo in nomes:
        if alvo in n:
            return n.index(alvo)
    return -1

def _data_cel(v):
    import datetime as _dt
    if v is None or v == "":
        return None
    if isinstance(v, (_dt.datetime, _dt.date)):
        return v.date() if isinstance(v, _dt.datetime) else v
    s = str(v).strip()[:10]
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return _dt.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None

def _data_br(s):
    import datetime as _dt
    try:
        return _dt.datetime.strptime(s.strip(), "%d/%m/%Y").date()
    except (ValueError, AttributeError):
        return None

def agregar_pedidos(ws, cab, de=None, ate=None):
    """Relatório de pedidos Shopee → produtos agregados + estatísticas de saúde."""
    i_prod = _idx(cab, ["nome do produto"])
    i_qtd = _idx(cab, ["quantidade"])
    i_sub = _idx(cab, ["subtotal do produto"])
    i_preco = _idx(cab, ["preco acordado", "preço acordado"])
    i_status = _idx(cab, ["status do pedido"])
    i_mot = _idx(cab, ["cancelar motivo"])
    i_data = _idx(cab, ["data de criacao do pedido", "data de criação do pedido"])
    i_total = _idx(cab, ["valor total"])
    i_txt = _idx(cab, ["taxa de transacao", "taxa de transação"])
    i_txc = _idx(cab, ["taxa de comissao liquida", "taxa de comissão líquida"])
    i_txs = _idx(cab, ["taxa de servico liquida", "taxa de serviço líquida"])

    def _f(v):
        try:
            return float(str(v).replace(",", ".")) if v not in (None, "") else 0.0
        except (ValueError, TypeError):
            return 0.0
    agg, det = {}, {"total": 0, "validos": 0, "cancelados": 0, "v_cancelado": 0.0,
                    "nao_pagos": 0, "andamento": 0, "v_risco": 0.0, "motivos": {},
                    "ignorados_data": 0, "por_produto": {}}
    for lin in ws.iter_rows(min_row=2, values_only=True):
        if not lin or all(c is None or str(c).strip() == "" for c in lin):
            continue
        det["total"] += 1
        if de or ate:
            d = _data_cel(lin[i_data] if 0 <= i_data < len(lin) else None)
            if d and ((de and d < de) or (ate and d > ate)):
                det["ignorados_data"] += 1
                continue
        st = str(lin[i_status] if 0 <= i_status < len(lin) and lin[i_status] is not None else "")
        sl = norm(st)
        ruim = ("cancelad" in sl) or ("não pago" in sl) or ("nao pago" in sl)
        if ruim and "cancelad" in sl:
            det["cancelados"] += 1
        elif ruim:
            det["nao_pagos"] += 1
        mot = ""
        if 0 <= i_mot < len(lin) and lin[i_mot] not in (None, ""):
            mot = str(lin[i_mot]).strip()
        try: qtd = int(float(str(lin[i_qtd]).replace(",", "."))) if i_qtd >= 0 else 0
        except (ValueError, TypeError): qtd = 0
        try: sub = float(str(lin[i_sub]).replace(",", ".")) if i_sub >= 0 else 0
        except (ValueError, TypeError): sub = 0.0
        if sub <= 0 and i_preco >= 0:
            try: sub = float(str(lin[i_preco]).replace(",", "."))*qtd
            except (ValueError, TypeError): pass
        prod = str(lin[i_prod]).strip() if i_prod >= 0 and lin[i_prod] else ""
        if ruim:
            if mot:
                det["motivos"][mot] = det["motivos"].get(mot, 0) + 1
            det["v_cancelado"] += sub
            if prod:
                pc = det["por_produto"].setdefault(prod, {"qtd": 0, "valor": 0.0, "motivos": {}})
                pc["qtd"] += qtd
                pc["valor"] = round(pc["valor"] + sub, 2)
                if mot:
                    pc["motivos"][mot] = pc["motivos"].get(mot, 0) + 1
            continue
        if de or ate:
            d = _data_cel(lin[i_data] if i_data >= 0 else None)
            if d and ((de and d < de) or (ate and d > ate)):
                det["ignorados_data"] += 1
                continue
        if not prod or qtd <= 0:
            continue
        det["validos"] += 1
        if any(k in sl for k in ("enviado", "entregue", "devolu", "a enviar")):
            det["andamento"] += 1
            det["v_risco"] += sub
        a = agg.setdefault(prod, {"produto": prod, "qtd": 0, "fat": 0.0, "est": 0, "custo": 0, "taxas": 0.0})
        a["qtd"] += qtd
        a["fat"] = round(a["fat"] + sub, 2)
        a["taxas"] = round(a["taxas"] + _f(lin[i_txt] if 0 <= i_txt < len(lin) else 0)
                           + _f(lin[i_txc] if 0 <= i_txc < len(lin) else 0)
                           + _f(lin[i_txs] if 0 <= i_txs < len(lin) else 0), 2)
    if not agg:
        raise ValueError("Não achei vendas válidas. Só há pedidos cancelados/não pagos no arquivo?")
    return list(agg.values()), det

ULTIMOS_PEDIDOS = {}

def ler_excel(caminho, de=None, ate=None):
    try:
        import openpyxl
    except ImportError:
        raise ValueError("Falta uma peça (openpyxl). Rode: pip install openpyxl")
    try:
        wb = openpyxl.load_workbook(caminho, data_only=True)
    except PermissionError:
        raise ValueError("O arquivo está aberto no Excel. Feche ele lá e tente de novo.")
    except Exception:
        try:  # às vezes é um CSV com nome .xlsx
            return ler_csv(caminho)
        except Exception:
            raise ValueError("Não consegui abrir. Confira se é um Excel válido (.xlsx) ou CSV.")
    ws = wb.active
    linhas = list(ws.iter_rows(values_only=True))
    if not linhas or not any(linhas[0]):
        raise ValueError("Planilha vazia ou sem títulos na primeira linha.")
    cab = [(str(c) if c is not None else "").strip() for c in linhas[0]]
    if eh_relatorio_pedidos(cab):
        produtos, det = agregar_pedidos(ws, cab, de, ate)
        ULTIMOS_PEDIDOS.clear()
        ULTIMOS_PEDIDOS.update(det)
        return produtos
    regs = []
    for lin in linhas[1:]:
        if not lin or all(c is None or str(c).strip() == "" for c in lin): continue
        regs.append({cab[i]: ("" if lin[i] is None else str(lin[i]).strip())
                     for i in range(min(len(cab), len(lin)))})
    return _montar_lista(regs)

def exemplo():
    d = [("Camisa Oversized Preta", 420, 29.90, 180), ("Calça Cargo Bege", 310, 89.90, 90),
         ("Tênis Street Branco", 150, 199.90, 40), ("Boné Aba Reta", 500, 39.90, 300),
         ("Moletom Canguru Cinza", 120, 119.90, 55), ("Meia Cano Alto Kit 3", 800, 24.90, 500),
         ("Jaqueta Corta Vento", 60, 159.90, 25), ("Cinto Couro Marrom", 200, 49.90, 120),
         ("Bermuda Tactel", 350, 44.90, 200), ("Camiseta Básica Branca", 900, 19.90, 600),
         ("Relógio Digital Sport", 45, 129.90, 15), ("Carteira Slim", 110, 59.90, 70)]
    return [{"produto": p, "qtd": q, "fat": round(q*pr, 2), "est": e, "custo": round(pr*0.45, 2)} for p, q, pr, e in d]

# ================= contas =================
# Taxas oficiais vigentes (fontes: Central do Vendedor Shopee 01/10/2026; ML 02/03/2026).
# Shopee CNPJ: % por faixa + fixo por item (frete grátis já embutido). CPF c/ +450 ped/90d: +R$3/item.
def taxa_shopee(preco, cpf_extra=False, frete_extra=0.0):
    if preco < 9:
        return preco*0.5 + frete_extra
    if preco <= 79.99: pct, fixa = 0.20, 4.50
    elif preco < 100: pct, fixa = 0.14, 16.0
    elif preco < 200: pct, fixa = 0.14, 20.0
    else: pct, fixa = 0.14, 26.0
    if cpf_extra: fixa += 3.0
    return preco*pct + fixa + frete_extra

# ML: comissão Clássico ~12% / Premium ~17% (varia por categoria; editável) +
# envio por peso e faixa (reputação verde, Full/Coleta/Agência). Faixa 100-199 = estimativa.
ENVIO_ML = {
    "leve":  [(49, 6.85), (79, 8.15), (100, 12.95), (200, 13.90), (float("inf"), 21.65)],
    "medio": [(49, 7.35), (79, 8.65), (100, 14.75), (200, 16.50), (float("inf"), 25.45)],
}
def envio_ml(preco, peso="leve"):
    for teto, valor in ENVIO_ML.get(peso, ENVIO_ML["leve"]):
        if preco < teto: return valor
    return ENVIO_ML["leve"][-1][1]

def taxa_ml(preco, premium=False, com_pct=0.0, peso="leve", flex=False):
    com = (com_pct or (17.0 if premium else 12.0))/100
    total = preco*com
    if flex and preco < 79:
        total += 7.75  # custo fixo Flex/próprio abaixo de R$79
    else:
        total += envio_ml(preco, peso)
    return total

def preco_ideal_shopee(custo, ta, mlucro, cpf_extra=False, frete_extra=0.0, imp=0.0, dev=0.0):
    """Por quanto anunciar na Shopee dados custo e lucro desejado (iteração nas faixas)."""
    if custo <= 0 or 1 - 0.20 - ta - mlucro - imp - dev <= 0.05: return 0.0
    p = custo/(1 - 0.20 - ta - mlucro - imp - dev) + 5
    for _ in range(5):
        pct = 0.20 if p <= 79.99 else 0.14
        fixa = taxa_shopee(p, cpf_extra, frete_extra) - p*pct
        base = 1 - pct - ta - mlucro - imp - dev
        if base <= 0.05: return 0.0
        p = (custo + fixa)/base
    return round(p, 2)

def preco_ideal_ml(custo, ta, mlucro, com=0.12, peso="leve", flex=False, frete_man=0.0, imp=0.0, dev=0.0):
    """Por quanto anunciar no ML dados custo e lucro desejado."""
    base = 1 - com - ta - mlucro - imp - dev
    if custo <= 0 or base <= 0.05: return 0.0
    p = custo/base + 15
    for _ in range(5):
        if peso == "manual": fr = frete_man
        elif flex and p < 79: fr = 7.75
        else: fr = envio_ml(p, peso)
        p = (custo + fr)/base
    return round(p, 2)

def roas_alvo(margem_contrib, gordura=0.15):
    """Cálculo de especialista: breakeven = 1/margem; alvo = breakeven x (1+gordura)."""
    if margem_contrib <= 0.02:
        return 0.0, 0.0
    be = 1/margem_contrib
    return round(be, 1), round(be*(1+gordura), 1)

COMISSAO_PADRAO = {"Shopee": 20.0, "Mercado Livre": 12.0}
CPC_PADRAO = {"Shopee": 0.45, "Mercado Livre": 0.85}
CONV_PADRAO = {"Shopee": 0.025, "Mercado Livre": 0.02}

def faixa_qtd(n):
    return "Menos de 20" if n < 20 else ("De 20 a 100" if n <= 100 else "Mais de 100")

def solucao_motivo(mot):
    t = (mot or "").lower()
    if "não pago" in t or "nao pago" in t or "pagamento" in t or "não pagou" in t:
        return "Ative Pix, responda dúvidas de preço no chat e lembre o cliente antes de expirar."
    if "não enviou" in t or "nao enviou" in t or "estoque" in t:
        return "Mantenha estoque atualizado e poste em até 24h; pause anúncio sem estoque."
    if "mais barato" in t or "preço" in t or "preco" in t:
        return "Reveja o preço x concorrentes (aba 6) ou monte kit com brinde."
    if "endere" in t or "entrega" in t or "área" in t or "area" in t or "risco" in t or "rastre" in t:
        return "Confirme o endereço no chat e use envio com rastreio."
    if "arrepend" in t or "desist" in t or "engano" in t or "errado" in t or "tamanho" in t or "cor" in t:
        return "Melhore fotos, medidas e descrição para alinhar a expectativa."
    if "prazo" in t or "demora" in t or "atraso" in t:
        return "Ajuste o prazo no anúncio e poste no mesmo dia."
    return "Leia o motivo e corrija o ponto no anúncio ou na operação."

def proxima_campanha():
    """Retorna (nome, dias_faltando, reserva_extra) da próxima data forte de vendas."""
    import datetime as _dt
    import calendar as _cal
    hoje = _dt.date.today()

    def seg_dom(ano, mes):
        ds = [s[6] for s in _cal.monthcalendar(ano, mes) if s[6]]
        return _dt.date(ano, mes, ds[1])

    def bf(ano):
        sx = [s[4] for s in _cal.monthcalendar(ano, 11) if s[4]]
        return _dt.date(ano, 11, sx[3])

    evs = []
    for ano in (hoje.year, hoje.year+1):
        for m, d, n in [(9, 9, "9.9"), (10, 10, "10.10"), (11, 11, "11.11"),
                        (12, 12, "12.12"), (12, 25, "Natal"), (1, 1, "Ano Novo"),
                        (3, 15, "Dia do Consumidor"), (10, 12, "Dia das Crianças")]:
            evs.append((_dt.date(ano, m, d), n))
        evs.append((seg_dom(ano, 5), "Dia das Mães"))
        evs.append((seg_dom(ano, 8), "Dia dos Pais"))
        evs.append((bf(ano), "Black Friday"))
    fut = sorted((dt, n) for dt, n in evs if dt >= hoje)
    if not fut:
        return "rotina", 30, 0.15
    dt, nome = fut[0]
    dias = (dt-hoje).days
    extra = 0.50 if dias <= 15 else (0.30 if dias <= 30 else 0.15)
    return nome, dias, extra

def classificar(dados):
    ordenados = sorted(dados, key=lambda x: x["fat"], reverse=True)
    total = sum(x["fat"] for x in ordenados) or 1
    acum, res = 0.0, []
    for x in ordenados:
        acum += x["fat"]
        pct = acum / total * 100
        classe = "A" if pct <= 80 else ("B" if pct <= 95 else "C")
        res.append({**x, "classe": classe, "acum": round(pct, 1),
                    "part": round(x["fat"]/total*100, 2)})
    return res, round(total, 2)

def dinheiro(v):
    return f"R$ {v:,.2f}".replace(",", ".")

AJUDA = {
    "curva": ("O que é a Curva ABC?\n\nÉ uma forma de descobrir quais produtos trazem mais dinheiro.\n"
              "• Classe A: os campeões, que juntos somam 80% do seu faturamento.\n"
              "• Classe B: os do meio, que podem virar campeões.\n"
              "• Classe C: os que vendem pouco e precisam de atenção."),
    "roas": ("O que é a meta de retorno?\n\nExemplo: meta 4x quer dizer que a cada R$ 1,00 gasto em anúncio, buscamos R$ 4,00 em vendas.\n"
             "O programa calcula sozinho a partir do seu lucro."),
    "verba": ("Como dividir a verba?\n\nOs campeões (classe A) merecem a maior parte.\n"
              "O padrão é 60% para A, 30% para B e 10% para C. A soma sempre precisa dar 100."),
    "grupo": ("Sozinho ou em kit?\n\n• Sozinho: o produto tem anúncio próprio. Ideal para campeões.\n"
              "• Em kit: o produto só é anunciado junto com um campeão. Ideal para produtos parados."),
}

# ================= interface =================
class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Central REH - Análise completa para o seu negócio")
        self.configure(bg=FUNDO)
        self.minsize(940, 640)
        self._centralizar()
        self.dados = []; self.resultado = []; self.total = 0
        self.loja = self._loja_load()
        self._estilo(); self._cabecalho(); self._abas(); self._rodape()
        self._status()

    # ----- dados salvos (estoque e custo ficam guardados) -----
    def _loja_path(self):
        return os.path.join(os.path.dirname(os.path.abspath(__file__)), "minha_loja.json")

    def _loja_load(self):
        try:
            with open(self._loja_path(), encoding="utf-8") as f:
                d = json.load(f)
            return d if isinstance(d, dict) else {}
        except (OSError, ValueError):
            return {}

    def _loja_save(self):
        try:
            with open(self._loja_path(), "w", encoding="utf-8") as f:
                json.dump(self.loja, f, ensure_ascii=False)
        except OSError:
            pass

    def _aplicar_loja(self):
        """Preenche estoque/custo salvos nos produtos importados (pelo nome)."""
        n = 0
        for x in self.dados:
            reg = self.loja.get(x["produto"])
            if reg:
                if x.get("est", 0) in (0, None) and reg.get("est"):
                    x["est"] = reg["est"]; n += 1
                if x.get("custo", 0) in (0, None) and reg.get("custo"):
                    x["custo"] = reg["custo"]; n += 1
        return n

    def _tela_real(self):
        """Tamanho do monitor principal (winfo_screenwidth soma todos os monitores e erra o centro)."""
        try:
            import ctypes, sys
            if sys.platform == "win32":
                u = ctypes.windll.user32
                return u.GetSystemMetrics(0), u.GetSystemMetrics(1)
        except Exception:
            pass
        return self.winfo_screenwidth(), self.winfo_screenheight()

    def _centralizar(self):
        self.update_idletasks()
        sw, sh = self._tela_real()
        w, h = min(940, sw - 40), min(660, sh - 60)
        self.geometry(f"{w}x{h}+{max(0, (sw-w)//2)}+{max(0, (sh-h)//2)}")
        self.update_idletasks()
        self.lift()
        self.attributes("-topmost", True)
        self.after(600, lambda: self.attributes("-topmost", False))
        self.after(1200, self._conter)
        self.focus_force()

    def _conter(self):
        """Traz a janela para dentro do monitor principal se o sistema a empurrou para fora."""
        try:
            self.update_idletasks()
            sw, sh = self._tela_real()
            w, h, x, y = self.winfo_width(), self.winfo_height(), self.winfo_x(), self.winfo_y()
            if w > sw - 20: w = sw - 20
            if h > sh - 60: h = sh - 60
            if x + w > sw - 10: x = max(0, sw - 10 - w)
            if x < 0: x = 0
            if y < 0: y = 0
            self.geometry(f"{w}x{h}+{x}+{y}")
        except Exception:
            pass

    def _estilo(self):
        st = ttk.Style(self); st.theme_use("clam")
        st.configure(".", font=("Segoe UI", 12), background=FUNDO, foreground=CINZA_TXT)
        st.configure("TButton", font=("Segoe UI", 12, "bold"), padding=(14, 9),
                     relief="flat", borderwidth=0, focusthickness=0, focuscolor="",
                     cursor="hand2", background="#EFE9DC")
        st.map("TButton", background=[("active", "#E2D8C2")])
        st.configure("Acao.TButton", background=LARANJA, foreground="#1A1A1A", padding=(14, 10),
                     font=("Segoe UI", 13, "bold"), relief="flat", borderwidth=0,
                     focusthickness=0, focuscolor="", cursor="hand2")
        st.map("Acao.TButton", background=[("active", "#A8821C"), ("disabled", "#D8CFB8")])
        st.configure("Ativo.TButton", background="#3D0A1A", foreground="white", padding=(14, 10),
                     font=("Segoe UI", 13, "bold"), relief="flat", borderwidth=0,
                     focusthickness=0, focuscolor="", cursor="hand2")
        st.map("Ativo.TButton", background=[("active", "#2E0712")])
        st.configure("TNotebook.Tab", font=("Segoe UI", 11, "bold"), padding=(6, 5))
        st.map("TNotebook.Tab", background=[("selected", "#F3E5C0")],
               foreground=[("selected", "#3D0A1A")])
        st.configure("TNotebook.Tab", font=("Segoe UI", 11, "bold"), padding=(6, 5))
        st.configure("Treeview", font=("Segoe UI", 12), rowheight=32)
        st.configure("Treeview.Heading", font=("Segoe UI", 12, "bold"), background=AZUL, foreground="white")
        st.configure("TEntry", padding=5, font=("Segoe UI", 12)); st.configure("TCombobox", padding=5)
        st.configure("TLabel", font=("Segoe UI", 12)); st.configure("TCheckbutton", font=("Segoe UI", 12))

    def _cabecalho(self):
        h = tk.Frame(self, bg=AZUL); h.pack(fill="x")
        box = tk.Frame(h, bg=AZUL); box.pack(side="left", padx=14, pady=8)
        tk.Label(box, text="Central REH", bg=AZUL, fg=CREME,
                 font=("Georgia", 22, "bold")).pack(anchor="w")
        tk.Label(box, text="Análise completa para o seu negócio", bg=AZUL, fg=OURO,
                 font=("Segoe UI", 11, "italic")).pack(anchor="w")
        self.passo = tk.Label(h, text="", bg=AZUL, fg=OURO,
                              font=("Segoe UI", 11, "bold"))
        self.passo.pack(side="right", padx=14)

    def _passo(self):
        try:
            i = self.nb.index(self.nb.select())
            nomes = ["Importar vendas", "Curva ABC", "Precificação", "Estratégia Ads",
                     "Oportunidades de melhoria", "Análise de concorrência", "Cancelados"]
            self.passo.config(text=f"Passo {i+1} de 7: {nomes[i]}")
        except Exception:
            pass
        try:
            if i == 5:
                self._conc_combo()
            if i == 6:
                self._canc_fill()
        except Exception:
            pass

    def _rodape(self):
        self.barra = tk.Label(self, text="", bg="#F3E5C0", fg=CINZA_TXT,
                              font=("Segoe UI", 11), anchor="w")
        self.barra.pack(side="bottom", fill="x", ipady=6)

    def _status(self):
        n = len(self.dados)
        self.barra.config(text=f"   Produtos: {n}    Análise: {'pronta' if self.resultado else 'pendente'}")

    def _abas(self):
        nb = ttk.Notebook(self); nb.pack(fill="both", expand=True, padx=10, pady=10)
        self.nb = nb
        self.tD = ttk.Frame(nb); self.tA = ttk.Frame(nb); self.tP = ttk.Frame(nb); self.tV = ttk.Frame(nb)
        self.tM = ttk.Frame(nb); self.tC = ttk.Frame(nb); self.tX = ttk.Frame(nb)
        nb.add(self.tD, text="  1 - Vendas  ")
        nb.add(self.tA, text="  2 - Curva ABC  ")
        nb.add(self.tP, text="  3 - Preços  ")
        nb.add(self.tV, text="  4 - Ads  ")
        nb.add(self.tM, text="  5 - Oportunidades  ")
        nb.add(self.tC, text="  6 - Concorrência  ")
        nb.add(self.tX, text="  7 - Cancelados  ")
        nb.bind("<<NotebookTabChanged>>", lambda e: self._passo())
        self.conc = {}
        self.slots = {1: None, 2: None, 3: None}
        self.slot_ativo = None
        self.arq_atual = ""
        self._tela_dados(); self._tela_analise(); self._tela_preco(); self._tela_verba()
        self._tela_melhorar(); self._tela_conc(); self._tela_canc()

    # ---------- peças visuais ----------
    def dica(self, master, texto, chave):
        f = tk.Frame(master, bg=ROSA_CLARO)
        f.pack(fill="x", padx=10, pady=(10, 4))
        tk.Label(f, text="DICA: " + texto, bg=ROSA_CLARO, fg=CINZA_TXT, wraplength=740,
                 justify="left", font=("Segoe UI", 12)).pack(side="left", padx=10, pady=8, fill="x", expand=True)
        ttk.Button(f, text="Explicar", width=10,
                   command=lambda: messagebox.showinfo("Explicando de forma simples", AJUDA[chave])).pack(side="right", padx=8)

    def barra_acoes(self, master, botoes):
        """botoes: lista de (texto, comando, destaque). Botões arredondados, 3 por fileira."""
        f = ttk.Frame(master); f.pack(fill="x", padx=10, pady=4)
        for i, (txt, cmd, dest) in enumerate(botoes):
            b = BotaoRedondo(f, txt, cmd, largura=290, altura=44,
                             fundo=LARANJA if dest else "#E7DFCF")
            b.grid(row=i // 3, column=i % 3, padx=4, pady=3, sticky="w")
        return f

    def tabela(self, master, colunas, altura):
        """colunas: [(id, titulo, largura, alinhar_esq)]. Com rolagem vertical."""
        box = ttk.Frame(master); box.pack(fill="both", expand=True, padx=10, pady=4)
        tv = ttk.Treeview(box, columns=[c[0] for c in colunas], show="headings", height=altura)
        for cid, tit, larg, esq in colunas:
            tv.heading(cid, text=tit)
            tv.column(cid, width=larg, minwidth=larg // 2, stretch=True, anchor="w" if esq else "center")
        sb = ttk.Scrollbar(box, orient="vertical", command=tv.yview)
        tv.configure(yscrollcommand=sb.set)
        tv.pack(side="left", fill="both", expand=True); sb.pack(side="right", fill="y")
        return tv

    # ---------- tela 1 ----------
    def _tela_dados(self):
        t = self.tD
        self.dica(t, "Mande sua planilha de vendas (Excel ou CSV, pode ser o relatório da Shopee). "
                     "Pode escolher VÁRIOS meses de uma vez (segure Ctrl). "
                     "Não precisa arrumar nada: a análise sai sozinha na aba 2.", "curva")
        self.barra_acoes(t, [
            ("Abrir minha planilha", self.imp, True),
            ("Saúde das vendas", self.saude_popup, False),
            ("Testar com exemplo", self.exe, False),
            ("Digitar um produto", self.add, False),
            ("Limpar tudo", self.limpar, False)])
        self.info = tk.Label(t, text="Nenhum produto ainda.",
                             bg=FUNDO, font=("Segoe UI", 13, "bold"))
        self.info.pack(anchor="w", padx=14, pady=2)
        ff = ttk.Frame(t); ff.pack(fill="x", padx=14, pady=2)
        ttk.Label(ff, text="Ver curva:", font=("Segoe UI", 11, "bold")).pack(side="left")
        self.filtro_val = "Todas"
        self.f_btns = {}
        for nome in ("Todas", "A", "B", "C"):
            b = ttk.Button(ff, text=nome if nome == "Todas" else f"Só {nome}", width=8,
                           command=lambda n=nome: self.f_set(n))
            b.pack(side="left", padx=3)
            self.f_btns[nome] = b
        self.f_btns["Todas"].config(style="Acao.TButton")
        pf = ttk.Frame(t); pf.pack(fill="x", padx=14, pady=2)
        ttk.Label(pf, text="Período dos dados: de").pack(side="left")
        self.p_de = ttk.Entry(pf, width=12); self.p_de.pack(side="left", padx=4)
        ttk.Label(pf, text="até").pack(side="left")
        self.p_ate = ttk.Entry(pf, width=12); self.p_ate.pack(side="left", padx=4)
        tk.Label(pf, text="(DD/MM/AAAA — detecto sozinho no relatório da Shopee)", bg=FUNDO, font=("Segoe UI", 9)).pack(side="left", padx=4)
        sl = ttk.Frame(t); sl.pack(fill="x", padx=14, pady=2)
        ttk.Label(sl, text="Comparar meses:").pack(side="left")
        self.slot_btns = {}
        self.slot_lbls = {}
        for _n in (1, 2, 3):
            b = ttk.Button(sl, text=f"Planilha {_n}", width=12,
                           command=lambda n=_n: self.slot_clicar(n))
            b.pack(side="left", padx=4)
            self.slot_btns[_n] = b
            lb = tk.Label(sl, text="vazia", bg=FUNDO, font=("Segoe UI", 9))
            lb.pack(side="left")
            self.slot_lbls[_n] = lb
        ttk.Button(sl, text="Comparar", command=self.slots_comparar).pack(side="left", padx=8)
        ttk.Button(sl, text="Limpar", command=self.slots_limpar).pack(side="left", padx=2)
        self.tree = self.tabela(t, [("produto", "PRODUTO", 480, True), ("qtd", "VENDEU", 90, False),
                                    ("fat", "FATUROU", 120, False), ("classe", "CLASSE", 80, False)], 8)
        self.tree.tag_configure("A", background="#DFF5DF"); self.tree.tag_configure("B", background="#FFF3D6")
        self.tree.tag_configure("C", background="#F0F0F0")
        nav = ttk.Frame(t); nav.pack(fill="x", padx=10, pady=8)
        BotaoRedondo(nav, "VER CURVA ABC  >>", lambda: self.nb.select(self.tA),
                     largura=880, altura=50).pack()

    def f_set(self, nome):
        if not self.resultado and self.dados:
            self.analisar()
        if not self.resultado:
            messagebox.showwarning("Calma lá", "Abra sua planilha primeiro (passo 1)."); return
        self.filtro_val = nome
        for n, b in self.f_btns.items():
            b.config(style="Acao.TButton" if n == nome else "TButton")
        self._refresh(True)

    # ---------- tela 2 ----------
    def _tela_analise(self):
        t = self.tA
        self.dica(t, "Aqui você descobre seus campeões (verde), quem pode crescer (amarelo) e quem está parado (cinza), "
                     "com o que fazer em cada caso.", "curva")
        self.barra_acoes(t, [
            ("Copiar texto", self.copiar, False),
            ("Salvar relatório", self.saverep, False),
            ("Salvar CSV", self.savecsv, False),
            ("Salvar Excel A/B/C", self.savexlsx, False),
            ("Comparar meses", self.hist_popup, False),
            ("Evolução", self.evo_popup, False),
            ("DRE real", self.dre_popup, False)])
        self.out = scrolledtext.ScrolledText(t, wrap="word", font=("Segoe UI", 12))
        self.out.pack(fill="both", expand=True, padx=10, pady=4)
        nav = ttk.Frame(t); nav.pack(fill="x", padx=10, pady=8)
        BotaoRedondo(nav, "MONTAR ESTRATÉGIA ADS  >>", lambda: self.nb.select(self.tV),
                     largura=880, altura=50).pack()

    # ---------- tela 3 ----------
    def _tela_verba(self):
        t = self.tV
        self.dica(t, "Diga quanto pode gastar por dia e onde vai anunciar. O programa monta sozinho: "
                     "quanto vai para cada produto, se anuncia sozinho ou em kit, e quanto deve voltar em vendas.", "grupo")
        grp = ttk.Frame(t); grp.pack(fill="x", padx=10, pady=4)
        for _c in range(3):
            grp.columnconfigure(_c, weight=1)
        g1 = ttk.LabelFrame(grp, text="1 - Onde anunciar", padding=8)
        g1.grid(row=0, column=0, sticky="nsew", padx=4)
        self.v_plat = ttk.Combobox(g1, values=["Shopee", "Mercado Livre", "Nos dois"], width=20, state="readonly")
        self.v_plat.set("Shopee"); self.v_plat.pack(fill="x", pady=2)
        self.v_mltipo = ttk.Combobox(g1, values=["Clássico (taxa menor)", "Premium (12x sem juros)"], width=20, state="readonly")
        self.v_mltipo.set("Clássico (taxa menor)"); self.v_mltipo.pack(fill="x", pady=2)
        g2 = ttk.LabelFrame(grp, text="2 - Verba e lucro", padding=8)
        g2.grid(row=0, column=1, sticky="nsew", padx=4)
        r1 = ttk.Frame(g2); r1.pack(fill="x", pady=2)
        ttk.Label(r1, text="Gasto/dia R$:").pack(side="left")
        self.v_total = ttk.Entry(r1, width=10); self.v_total.pack(side="left", padx=4)
        r2 = ttk.Frame(g2); r2.pack(fill="x", pady=2)
        ttk.Label(r2, text="Lucro líq. %:").pack(side="left")
        self.v_margem = ttk.Entry(r2, width=10); self.v_margem.insert(0, "22"); self.v_margem.pack(side="left", padx=4)
        ttk.Button(r2, text="?", width=3,
                   command=lambda: messagebox.showinfo("Explicando de forma simples", AJUDA["roas"])).pack(side="left")
        g3 = ttk.LabelFrame(grp, text="3 - Divisão A/B/C (%)", padding=8)
        g3.grid(row=0, column=2, sticky="nsew", padx=4)
        dv = ttk.Frame(g3); dv.pack(fill="x", pady=2)
        self.v_pa = ttk.Entry(dv, width=6); self.v_pa.insert(0, "60"); self.v_pa.pack(side="left")
        self.v_pb = ttk.Entry(dv, width=6); self.v_pb.insert(0, "30"); self.v_pb.pack(side="left", padx=4)
        self.v_pc = ttk.Entry(dv, width=6); self.v_pc.insert(0, "10"); self.v_pc.pack(side="left")
        tk.Label(dv, text="= 100", bg=FUNDO).pack(side="left", padx=4)
        ttk.Button(dv, text="?", width=3,
                   command=lambda: messagebox.showinfo("Explicando de forma simples", AJUDA["verba"])).pack(side="left")
        go = ttk.Frame(t); go.pack(fill="x", padx=10, pady=4)
        BotaoRedondo(go, "MONTAR MINHA ESTRATÉGIA", self.v_gerar,
                     largura=880, altura=50, fonte=("Segoe UI", 13, "bold")).pack()
        self.vtree = self.tabela(t, [("produto", "PRODUTO", 320, True), ("plat", "ONDE", 70, False),
                                      ("formato", "SOZINHO OU KIT?", 170, True), ("verba", "GASTAR/DIA", 100, False),
                                      ("roas", "META", 75, False), ("retorno", "DEVE VOLTAR", 115, False)], 6)
        self.v_out = scrolledtext.ScrolledText(t, wrap="word", height=5, font=("Segoe UI", 12))
        self.v_out.pack(fill="x", padx=10, pady=(0, 8))

    # ----- melhorar (prioridades) -----
    def _tela_melhorar(self):
        t = self.tM
        self.dica(t, "Sua fila de trabalho: o que arrumar primeiro em cada anúncio, do mais urgente ao menos urgente. "
                     "Quanto maior o impacto no seu bolso, mais alto na lista.", "curva")
        top = ttk.Frame(t); top.pack(fill="x", padx=10, pady=4)
        ttk.Button(top, text="Gerar prioridades", style="Acao.TButton", command=self.m_gerar).pack(side="left", padx=2)
        ttk.Button(top, text="Montar kit", command=self.kit_popup).pack(side="left", padx=2)
        ttk.Button(top, text="Informar estoque", command=self.estoque_popup).pack(side="left", padx=2)
        ttk.Button(top, text="Acompanhar ruptura", command=self.ruptura_popup).pack(side="left", padx=2)
        ttk.Button(top, text="Campanhas", command=self.camp_popup).pack(side="left", padx=2)
        ttk.Button(top, text="Foco C", command=self.fococ_popup).pack(side="left", padx=2)
        ttk.Button(top, text="Enviar p/ Full", command=self.full_popup).pack(side="left", padx=2)
        ttk.Button(top, text="Copiar lista", command=lambda: (self.clipboard_clear(), self.clipboard_append(self.m_out.get("1.0", "end")), self.update())).pack(side="left", padx=2)
        self.mtree = self.tabela(t, [("produto", "PRODUTO", 320, True), ("cl", "CL", 45, False),
                                     ("problema", "PROBLEMA", 230, True), ("acao", "O QUE FAZER", 280, True),
                                     ("impacto", "IMPACTO", 85, False)], 11)
        self.mtree.tag_configure("A", background="#DFF5DF"); self.mtree.tag_configure("B", background="#FFF3D6")
        self.mtree.tag_configure("C", background="#F0F0F0")
        self.m_out = scrolledtext.ScrolledText(t, wrap="word", height=5, font=("Segoe UI", 12))
        self.m_out.pack(fill="x", padx=10, pady=(0, 8))

    def m_gerar(self):
        if not self.dados:
            messagebox.showwarning("Calma lá", "Primeiro abra sua planilha no passo 1."); return
        if not self.resultado:
            self.resultado, self.total = classificar(self.dados); self._refresh(com_classe=True)
        dias = self._dias_periodo()
        precos = sorted(x["fat"]/x["qtd"] for x in self.resultado if x["qtd"])
        med = precos[len(precos)//2] if precos else 0
        rows = []
        for x in self.resultado:
            preco = x["fat"]/x["qtd"] if x["qtd"] else 0
            peso = x["part"]  # % do faturamento = impacto
            cand = []
            if not x.get("custo"):
                cand.append(("Lucro desconhecido (sem custo)", "Informar o custo no passo 3", 3 if x["classe"] == "A" else 2))
            elif preco:
                mg = (preco - x["custo"])/preco
                if mg < 0.15:
                    cand.append((f"Margem apertada ({mg*100:.0f}%)", "Subir preço ou trocar fornecedor", 3))
            if len(x["produto"]) < 40:
                cand.append(("Título curto p/ Shopee", "Estender p/ 80+ letras (modelo, medida, uso)", 2))
            if x["classe"] == "A" and 0 < x["est"] < x["qtd"]*0.3:
                cand.append(("Campeão com pouco estoque", "Repor urgente antes de anunciar", 3))
            if x.get("est", 0) > 0 and x["qtd"] > 0:
                cob = x["est"]/(x["qtd"]/dias)
                if cob < 14:
                    cand.append((f"Acaba em {cob:.0f} dias", f"Repor para {14* x['qtd']/dias:.0f} un (14 dias)".replace(",", "."),
                                 3 if x["classe"] == "A" else 2))
            if x["classe"] == "C" and x["est"] > max(x["qtd"]*2, 1):
                cand.append(("Parado com estoque alto", "Criar kit com campeão (10-15% off)", 2))
            if med and preco > med*1.3:
                cand.append(("Preço acima da sua média", "Justificar na foto/título ou ajustar", 2))
            if x["classe"] == "B":
                cand.append(("Pode virar campeão", "Testar anúncio sozinho no passo 4", 1))
            for prob, acao, sev in cand:
                rows.append((peso*sev, x["produto"], x["classe"], prob, acao, peso))
        rows.sort(reverse=True)
        for i in self.mtree.get_children(): self.mtree.delete(i)
        for _s, prod, cl, prob, acao, peso in rows[:25]:
            self.mtree.insert("", "end", values=(prod[:60], cl, prob, acao, f"{peso:.1f}%"), tags=(cl,))
        L = [f"Encontrei {len(rows)} pontos de melhoria em {len(self.resultado)} produtos.",
             "Comece de cima para baixo: o impacto mostra quanto do seu faturamento cada item representa.",
             "\nDepois de arrumar, cadastre os concorrentes desse produto na aba 6 para comparar preço e oferta."]
        self.m_out.delete("1.0", "end"); self.m_out.insert("end", "\n".join(L))
        self._status()

    # ----- concorrentes -----
    def _tela_conc(self):
        t = self.tC
        self.dica(t, "Escolha seu produto e clique em 'Buscar sozinho na internet' que eu preencho os concorrentes. "
                     "Depois eu comparo preço, vendas e oferta e digo o que fazer.", "grupo")
        f = ttk.Frame(t); f.pack(fill="x", padx=14, pady=4)
        f.columnconfigure(1, weight=1); f.columnconfigure(3, weight=1)
        ttk.Label(f, text="Produto meu:").grid(row=0, column=0, sticky="w", pady=3)
        self.c_prod = ttk.Combobox(f, state="readonly", width=40)
        self.c_prod.grid(row=0, column=1, sticky="we", padx=6, pady=3)
        self.c_prod.bind("<<ComboboxSelected>>", lambda e: self.c_sugere())
        ttk.Label(f, text="Minhas fotos (qtd):").grid(row=0, column=2, sticky="w", pady=3)
        self.c_minhas = ttk.Entry(f, width=8); self.c_minhas.grid(row=0, column=3, sticky="w", padx=6, pady=3)
        ttk.Label(f, text="Buscar por:").grid(row=1, column=0, sticky="w", pady=3)
        self.c_termo = ttk.Entry(f, width=40)
        self.c_termo.grid(row=1, column=1, sticky="we", padx=6, pady=3)
        tk.Label(f, text="(ajuste as palavras se vier coisa errada)", bg=FUNDO, font=("Segoe UI", 9)).grid(row=1, column=2, columnspan=2, sticky="w", pady=3)
        g = ttk.Frame(t); g.pack(fill="x", padx=14, pady=2)
        ttk.Label(g, text="Concorrente:").pack(side="left")
        self.c_nome = ttk.Entry(g, width=24); self.c_nome.pack(side="left", padx=4)
        ttk.Label(g, text="Preço:").pack(side="left")
        self.c_preco = ttk.Entry(g, width=9); self.c_preco.pack(side="left", padx=2)
        ttk.Label(g, text="Vend/mês:").pack(side="left")
        self.c_vend = ttk.Entry(g, width=7); self.c_vend.pack(side="left", padx=2)
        ttk.Label(g, text="Nota:").pack(side="left")
        self.c_nota = ttk.Entry(g, width=5); self.c_nota.pack(side="left", padx=2)
        ttk.Label(g, text="Fotos:").pack(side="left")
        self.c_fotos = ttk.Entry(g, width=5); self.c_fotos.pack(side="left", padx=2)
        ttk.Button(g, text="+ Add", width=6, command=self.c_add).pack(side="left", padx=4)
        g1 = ttk.Frame(t); g1.pack(fill="x", padx=14, pady=2)
        ttk.Label(g1, text="Link do anúncio:").pack(side="left")
        self.c_link = ttk.Entry(g1, width=70); self.c_link.pack(side="left", padx=4, fill="x", expand=True)
        tk.Label(g1, text="(duplo-clique na linha abre o anúncio)", bg=FUNDO, font=("Segoe UI", 9)).pack(side="left", padx=4)
        g2 = ttk.Frame(t); g2.pack(fill="x", padx=14, pady=2)
        ttk.Label(g2, text="Achar concorrentes:").pack(side="left")
        ttk.Button(g2, text="Abrir busca Shopee", command=self.c_busca_shopee).pack(side="left", padx=4)
        ttk.Button(g2, text="Abrir busca ML", command=self.c_busca_ml).pack(side="left", padx=2)
        ttk.Label(g2, text="Cole vários (1 por linha: nome | preço | vend | nota | fotos):").pack(side="left", padx=(8, 0))
        ttk.Button(g2, text="Importar colados", command=self.c_importar).pack(side="left", padx=4)
        g3 = ttk.Frame(t); g3.pack(fill="x", padx=14, pady=2)
        ttk.Button(g3, text="Buscar sozinho na internet", style="Acao.TButton", command=self.c_auto).pack(side="left", padx=2)
        ttk.Button(g3, text="Chave API", command=self.c_chave).pack(side="left", padx=2)
        ttk.Button(g3, text="Apagar lista", command=self.c_limpar).pack(side="left", padx=2)
        tk.Label(g3, text="Automático usa Google Shopping (grátis). Traz nome, preço e loja; as vendas/mês você confere no anúncio.",
                 bg=FUNDO, font=("Segoe UI", 9)).pack(side="left", padx=8)
        self.c_cola = tk.Text(t, height=3, wrap="word", font=("Segoe UI", 10))
        self.c_cola.pack(fill="x", padx=14, pady=2)
        ttk.Button(t, text="Analisar (foco nos 4 mais vendidos)", style="Acao.TButton", command=self.c_analisar).pack(fill="x", padx=10, pady=2)
        self.ctree = self.tabela(t, [("nome", "CONCORRENTE", 300, True), ("preco", "PREÇO", 90, False),
                                     ("vend", "VEND/MÊS", 85, False), ("nota", "NOTA", 70, False),
                                     ("fotos", "FOTOS", 65, False), ("site", "LINK", 70, False)], 4)
        self.ctree.bind("<Double-1>", lambda e: self.c_abrir())
        self.c_out = scrolledtext.ScrolledText(t, wrap="word", height=5, font=("Segoe UI", 12))
        self.c_out.pack(fill="x", padx=10, pady=(0, 8))

    def _termo_prod(self):
        manual = self.c_termo.get().strip()
        if manual:
            return manual
        prod = self.c_prod.get()
        if not prod:
            messagebox.showwarning("Selecione", "Escolha seu produto primeiro."); return None
        return self._sugestao(prod)

    @staticmethod
    def _sugestao(prod):
        palavras = [p for p in prod.split() if len(p) > 1 and p.lower() not in ("de", "da", "do", "e", "com", "+", "-", "para")]
        return " ".join(palavras[:8])

    def c_sugere(self):
        prod = self.c_prod.get()
        if not prod: return
        self.c_termo.delete(0, "end")
        self.c_termo.insert(0, self._sugestao(prod))

    def c_busca_shopee(self):
        termo = self._termo_prod()
        if termo: webbrowser.open(f"https://www.shopee.com.br/search?keyword={quote_plus(termo)}")

    def c_busca_ml(self):
        termo = self._termo_prod()
        if termo: webbrowser.open(f"https://lista.mercadolivre.com.br/{quote_plus(termo.replace(' ', '-'))}")

    def c_importar(self):
        prod = self.c_prod.get()
        if not prod:
            messagebox.showwarning("Selecione", "Escolha seu produto primeiro."); return
        n = 0
        for ln in self.c_cola.get("1.0", "end").splitlines():
            for sep in ("|", ";", "\t"):
                if sep in ln:
                    p = [x.strip() for x in ln.split(sep)]
                    break
            else:
                continue
            if len(p) < 2 or not p[0]: continue
            try:
                self.conc.setdefault(prod, []).append({
                    "nome": p[0][:60], "preco": num(p[1]),
                    "vend": int(num(p[2])) if len(p) > 2 and p[2] else 0,
                    "nota": num(p[3]) if len(p) > 3 and p[3] else 0.0,
                    "fotos": int(num(p[4])) if len(p) > 4 and p[4] else 0,
                    "link": p[5].strip() if len(p) > 5 else ""})
                n += 1
            except ValueError:
                continue
        self.c_cola.delete("1.0", "end")
        self._conc_refresh(prod)
        messagebox.showinfo("Pronto", f"{n} concorrentes importados!")

    def _conc_combo(self):
        base = [x["produto"] for x in (self.resultado or self.dados)]
        self.c_prod.config(values=base)
        if base and not self.c_prod.get():
            self.c_prod.set(base[0])
            self.c_sugere()

    def c_add(self):
        prod = self.c_prod.get()
        if not prod:
            messagebox.showwarning("Selecione", "Escolha seu produto primeiro."); return
        try:
            nome = self.c_nome.get().strip() or "Concorrente"
            c = {"nome": nome, "preco": num(self.c_preco.get()), "vend": int(num(self.c_vend.get())),
                 "nota": num(self.c_nota.get()), "fotos": int(num(self.c_fotos.get())),
                 "link": self.c_link.get().strip()}
        except ValueError:
            messagebox.showerror("Ops", "Use só números em preço, vendas, nota e fotos."); return
        if c["preco"] <= 0:
            messagebox.showwarning("Faltou o preço", "Informe ao menos o preço do concorrente."); return
        self.conc.setdefault(prod, []).append(c)
        self._conc_refresh(prod)
        for e in (self.c_nome, self.c_preco, self.c_vend, self.c_nota, self.c_fotos, self.c_link):
            e.delete(0, "end")

    def _conc_refresh(self, prod):
        for i in self.ctree.get_children(): self.ctree.delete(i)
        for c in self.conc.get(prod, []):
            if "link" not in c:
                c["link"] = ""
            self.ctree.insert("", "end", values=(c["nome"][:50], dinheiro(c["preco"]), c["vend"],
                                                 f"{c['nota']:.1f}" if c["nota"] else "—",
                                                 c["fotos"] or "—",
                                                 "abrir >" if c["link"] else "—"))

    def c_abrir(self):
        prod = self.c_prod.get()
        sel = self.ctree.selection()
        if not prod or not sel:
            return
        idx = self.ctree.index(sel[0])
        conc = self.conc.get(prod, [])
        if 0 <= idx < len(conc) and conc[idx].get("link"):
            webbrowser.open(conc[idx]["link"])
        else:
            messagebox.showinfo("Sem link", "Esse concorrente não tem link.\nCole no campo 'Link do anúncio' e clique + Add de novo.")

    def c_limpar(self):
        prod = self.c_prod.get()
        if prod and prod in self.conc:
            self.conc[prod] = []
            self._conc_refresh(prod)

    def _chave_path(self):
        import os
        return os.path.join(os.path.dirname(os.path.abspath(__file__)), "chave_serpapi.txt")

    def c_chave(self):
        atual = ""
        try:
            with open(self._chave_path(), encoding="utf-8") as f:
                atual = f.read().strip()
        except OSError:
            atual = ""
        nova = simpledialog.askstring("Chave API",
            "Cole sua chave grátis do SerpApi (crie em serpapi.com → Register → Dashboard → API Key).\n"
            f"{'Você já tem uma salva.' if atual else 'Sem chave salva.'}\nDeixe vazio para apagar.")
        if nova is None: return
        try:
            with open(self._chave_path(), "w", encoding="utf-8") as f:
                f.write(nova.strip())
            messagebox.showinfo("Pronto", "Chave salva!" if nova.strip() else "Chave apagada.")
        except OSError as e:
            messagebox.showerror("Ops", f"Não consegui salvar: {e}")

    def c_auto(self):
        prod = self.c_prod.get()
        if not prod:
            messagebox.showwarning("Selecione", "Escolha seu produto primeiro."); return
        try:
            with open(self._chave_path(), encoding="utf-8") as f:
                key = f.read().strip()
        except OSError:
            key = ""
        if not key:
            messagebox.showwarning("Falta a chave",
                "Para buscar sozinho, preciso da chave grátis do SerpApi.\n\n1. Abra serpapi.com e crie sua conta (grátis, 250 buscas/mês).\n"
                "2. Copie a API Key do painel.\n3. Clique em 'Chave API' aqui e cole.\n\nDepois clique em 'Buscar sozinho' de novo.")
            return
        termo = self._termo_prod()
        if not termo: return
        url = (f"https://serpapi.com/search.json?engine=google_shopping&q={quote_plus(termo)}"
               f"&gl=br&hl=pt-BR&num=10&api_key={quote_plus(key)}")
        try:
            data = json.loads(http_get(url, timeout=20))
        except urllib.error.HTTPError as e:
            try:
                detalhe = json.loads(e.read().decode("utf-8", "ignore")).get("error", "")
            except Exception:
                detalhe = ""
            if e.code in (401, 403):
                messagebox.showerror("Chave rejeitada",
                    "O serviço não aceitou a chave. Confira se copiou ela inteira em 'Chave API'.\n" + detalhe)
            else:
                messagebox.showerror("Erro na busca", f"O serviço respondeu HTTP {e.code}. {detalhe}\nTente de novo.")
            return
        except Exception as e:
            messagebox.showerror("Sem internet?", f"Não consegui buscar agora ({e}). Tente de novo ou use manual."); return
        if data.get("error"):
            messagebox.showerror("Chave inválida?", f"O serviço respondeu: {data['error']}\nConfira a chave em 'Chave API'."); return
        achados = self._parse_shopping(data)
        if not achados:
            messagebox.showwarning("Nada achado", "Não achei ofertas para esse produto. Tente o manual."); return
        self.conc.setdefault(prod, []).extend(achados)
        self._conc_refresh(prod)
        messagebox.showinfo("Preenchido!",
            f"Busquei por '{termo}' e adicionei {len(achados)} concorrentes com nome, preço e loja.\n\nFalta só as VENDAS/mês (o Google não mostra): "
            "olhe os 4 mais vendidos e digite na linha deles com '+ Add' usando o mesmo nome.")

    @staticmethod
    def _parse_shopping(data):
        out = []
        for r in (data.get("shopping_results") or [])[:8]:
            try:
                preco = float(r.get("extracted_price") or 0)
            except (ValueError, TypeError):
                preco = 0.0
            if not preco:
                preco = num(r.get("price", ""))
            if preco <= 0: continue
            loja = (r.get("source") or "").strip()
            nome = (r.get("title") or "").strip()[:55]
            if loja: nome = f"{nome} ({loja})"
            try: nota = float(r.get("rating") or 0)
            except (ValueError, TypeError): nota = 0.0
            out.append({"nome": nome, "preco": round(preco, 2), "vend": 0, "nota": nota, "fotos": 0,
                        "link": r.get("link", "") or ""})
        return out

    def c_analisar(self):
        prod = self.c_prod.get()
        todos = self.conc.get(prod, [])
        if not prod or not todos:
            messagebox.showwarning("Faltam dados", "Escolha o produto e adicione ao menos 1 concorrente."); return
        conc = sorted(todos, key=lambda c: c["vend"], reverse=True)[:4]  # foco nos 4 mais vendidos
        meus = [x for x in self.dados if x["produto"] == prod]
        if not meus:
            messagebox.showwarning("Ops", "Não achei esse produto nos seus dados."); return
        meu = meus[0]
        pm = meu["fat"]/meu["qtd"] if meu["qtd"] else 0
        vals = sorted(c["preco"] for c in conc)
        vmin, vmax = vals[0], vals[-1]
        med = vals[len(vals)//2]
        vend_c = sum(c["vend"] for c in conc)
        L = [f"VOCÊ x TOP {len(conc)} CONCORRENTES (os mais vendidos) — {prod[:60]}", "="*60,
             f"Seu preço: {dinheiro(pm)} | Top concorrentes: {len(conc)} (de {dinheiro(vmin)} a {dinheiro(vmax)}, meio {dinheiro(med)})",
             f"Suas vendas: {meu['qtd']} un | Somados, eles vendem ~{vend_c} un/mês"]
        if pm > med:
            L.append(f"\nVocê está {((pm-med)/med*100):.0f}% ACIMA do meio. Para reagir: iguale em {dinheiro(med*0.97)} ou monte kit com brinde mantendo o preço.")
        elif pm < vmin:
            L.append(f"\nVocê é o MAIS BARATO ({((vmin-pm)/vmin*100):.0f}% abaixo). Cuidado com margem: teste subir para {dinheiro(vmin*0.98)} e observe as vendas.")
        else:
            L.append("\nSeu preço está na faixa. A briga agora é por foto, título e avaliação.")
        try: mf = int(self.c_minhas.get() or 0)
        except ValueError: mf = 0
        fotos_c = [c["fotos"] for c in conc if c["fotos"]]
        if mf and fotos_c and mf < max(fotos_c):
            L.append(f"Suas fotos ({mf}) < melhor concorrente ({max(fotos_c)}): adicione medida com régua, em uso e conteúdo da caixa.")
        notas = [c["nota"] for c in conc if c["nota"]]
        if notas:
            L.append(f"Nota média deles: {sum(notas)/len(notas):.1f}. Leia as avaliações 1-3 estrelas deles e corrija esses pontos no seu anúncio.")
        alvo = min(med*0.97, pm) if pm > med else pm
        L.append(f"\nPreço sugerido: {dinheiro(alvo)} + título de 80+ letras + 7 fotos. Revise em 7 dias.")
        self.c_out.delete("1.0", "end"); self.c_out.insert("end", "\n".join(L))

    # ----- saúde das vendas -----
    def saude_popup(self):
        s = getattr(self, "saude", None) or {}
        if not s:
            messagebox.showinfo("Saúde das vendas",
                "Importe o relatório de PEDIDOS da Shopee (Order...xlsx) para ver cancelamentos, motivos e valores.\n"
                "O relatório de produtos (Product_Sales) não tem esses detalhes."); return
        tot = s["total"] or 1
        tx_canc = s["cancelados"]/tot*100
        L = ["SAÚDE DAS VENDAS — explicada de forma simples", self._periodo_txt(), "="*60,
             f"Pedidos no período: {s['total']}",
             f"Vendas válidas: {s['validos']}  |  Em andamento/entregue: {s['andamento']} ({dinheiro(s['v_risco'])} ainda em risco)",
             f"Cancelados: {s['cancelados']} ({tx_canc:.1f}%) = {dinheiro(s['v_cancelado'])} perdidos",
             f"Não pagos: {s['nao_pagos']}"]
        if s["ignorados_data"]:
            L.append(f"(Fora do período escolhido: {s['ignorados_data']} ignorados.)")
        mots = sorted(s["motivos"].items(), key=lambda x: x[1], reverse=True)[:5]
        if mots:
            L.append("\nMOTIVOS DE CANCELAMENTO (do mais comum):")
            for mot, q in mots:
                L.append(f"  • {q}x — {mot[:90]}")
        L.append("\nO QUE MELHORAR:")
        if tx_canc >= 15:
            L.append("• Cancelamento alto! Confira estoque (venda sem estoque gera cancela), prazo de envio e descrição enganosa.")
        elif tx_canc >= 5:
            L.append("• Cancelamento moderado. Olhe o motivo nº 1 acima e corrija no anúncio.")
        else:
            L.append("• Cancelamento sob controle. Mantenha estoque e responda rápido.")
        if any("não pago" in m.lower() or "pagamento" in m.lower() for m in s["motivos"]):
            L.append("• Muito 'não pago': ative Pix e responda dúvidas de preço no chat.")
        if s["v_risco"] > 0:
            L.append(f"• {dinheiro(s['v_risco'])} ainda podem voltar (janela de devolução): capriche na pós-venda.")
        L.append("• Leia as 3 reclamações mais recentes da loja e corrija o ponto repetido.")
        pop = tk.Toplevel(self)
        pop.title("Saúde das vendas")
        pop.geometry("640x560")
        out = scrolledtext.ScrolledText(pop, wrap="word", font=("Segoe UI", 12))
        out.pack(fill="both", expand=True, padx=10, pady=10)
        out.insert("end", "\n".join(L))
        out.config(state="disabled")

    # ----- histórico mês a mês -----
    def _hist_path(self):
        return os.path.join(os.path.dirname(os.path.abspath(__file__)), "historico.json")

    def _hist_load(self):
        try:
            with open(self._hist_path(), encoding="utf-8") as f:
                d = json.load(f)
            return d if isinstance(d, dict) else {}
        except (OSError, ValueError):
            return {}

    def _hist_save(self):
        if not self.resultado:
            return
        per = f"{self.p_de.get().strip()} a {self.p_ate.get().strip()}".strip(" a") or "sem período"
        if per == "sem período":
            import datetime as _dt
            per = "importado em " + _dt.date.today().strftime("%d/%m/%Y")
        try:
            h = self._hist_load()
            saude_mini = {}
            if getattr(self, "saude", None):
                saude_mini = {"cancelados": self.saude.get("cancelados", 0),
                              "v_cancelado": round(self.saude.get("v_cancelado", 0.0), 2),
                              "total": self.saude.get("total", 0)}
            h[per] = {"total": self.total, "saude": saude_mini,
                      "produtos": {x["produto"]: {"qtd": x["qtd"], "fat": x["fat"], "classe": x["classe"]}
                                   for x in self.resultado}}
            with open(self._hist_path(), "w", encoding="utf-8") as f:
                json.dump(h, f, ensure_ascii=False)
        except OSError:
            pass

    def hist_popup(self):
        h = self._hist_load()
        if len(h) < 2:
            messagebox.showinfo("Histórico",
                "Preciso de ao menos 2 meses salvos para comparar.\nA cada análise com período, eu salvo sozinho."); return
        pop = tk.Toplevel(self)
        pop.title("Comparar meses")
        pop.geometry("700x640")
        f = ttk.Frame(pop); f.pack(fill="x", padx=10, pady=8)
        pers = sorted(h.keys())
        ttk.Label(f, text="De:").pack(side="left")
        c1 = ttk.Combobox(f, values=pers, width=22, state="readonly"); c1.set(pers[-2]); c1.pack(side="left", padx=4)
        ttk.Label(f, text="Para:").pack(side="left")
        c2 = ttk.Combobox(f, values=pers, width=22, state="readonly"); c2.set(pers[-1]); c2.pack(side="left", padx=4)
        chips = tk.Frame(pop, bg=FUNDO); chips.pack(fill="x", padx=10, pady=(0, 4))
        chip_a = tk.Label(chips, text="", bg="#3D0A1A", fg="white", font=("Segoe UI", 10, "bold"),
                          padx=10, pady=3)
        chip_a.pack(side="left", padx=(0, 6))
        chip_b = tk.Label(chips, text="", bg="#C9A227", fg="#1A1A1A", font=("Segoe UI", 10, "bold"),
                          padx=10, pady=3)
        chip_b.pack(side="left")
        banner = tk.Label(pop, text="", bg="#DFF5DF", fg="#1E7E34", font=("Segoe UI", 13, "bold"),
                          padx=10, pady=8, wraplength=660, justify="left")
        banner.pack(fill="x", padx=10, pady=(0, 4))
        tv = self.tabela(pop, [("m", "MEDIDA", 200, True),
                                 ("pa", c1.get()[:18], 150, False),
                                 ("pb", c2.get()[:18], 150, False)], 6)
        det = tk.Label(pop, text="", bg=FUNDO, font=("Segoe UI", 11, "bold"), anchor="w", justify="left")
        det.pack(fill="x", padx=10, pady=(0, 2))
        rol = tk.Canvas(pop, bg=FUNDO, highlightthickness=0)
        sb = ttk.Scrollbar(pop, orient="vertical", command=rol.yview)
        rol.configure(yscrollcommand=sb.set)
        sb.pack(side="right", fill="y")
        rol.pack(side="left", fill="both", expand=True, padx=(10, 0))
        corpo = ttk.Frame(rol)
        rol.create_window((0, 0), window=corpo, anchor="nw")
        corpo.bind("<Configure>", lambda e: rol.configure(scrollregion=rol.bbox("all")))
        rol.bind("<Enter>", lambda e: rol.bind_all("<MouseWheel>", lambda ev: rol.yview_scroll(-1*(ev.delta//120), "units")))
        rol.bind("<Leave>", lambda e: rol.unbind_all("<MouseWheel>"))
        ttk.Label(corpo, text="Quem SUBIU de curva", font=("Segoe UI", 11, "bold"),
                  foreground="#1E7E34").grid(row=0, column=0, sticky="w", pady=(4, 0))
        ttk.Label(corpo, text="Quem CAIU de curva", font=("Segoe UI", 11, "bold"),
                  foreground="#B71C1C").grid(row=0, column=1, sticky="w", pady=(4, 0))
        tsub = ttk.Treeview(corpo, columns=("p", "m"), show="headings", height=5)
        tsub.heading("p", text="PRODUTO"); tsub.heading("m", text="FOI PARA")
        tsub.column("p", width=230, anchor="w"); tsub.column("m", width=70, anchor="center")
        tsub.tag_configure("UP", background="#DFF5DF")
        tsub.grid(row=1, column=0, sticky="nsew", padx=(0, 4))
        tcai = ttk.Treeview(corpo, columns=("p", "m"), show="headings", height=5)
        tcai.heading("p", text="PRODUTO"); tcai.heading("m", text="CAIU PARA")
        tcai.column("p", width=230, anchor="w"); tcai.column("m", width=70, anchor="center")
        tcai.tag_configure("DN", background="#FFE3E3")
        tcai.grid(row=1, column=1, sticky="nsew", padx=(4, 0))
        corpo.columnconfigure(0, weight=1); corpo.columnconfigure(1, weight=1)
        ttk.Label(corpo, text="Maiores altas e quedas (R$)", font=("Segoe UI", 11, "bold")).grid(
            row=2, column=0, columnspan=2, sticky="w", pady=(6, 0))
        tmov = ttk.Treeview(corpo, columns=("p", "v"), show="headings", height=6)
        tmov.heading("p", text="PRODUTO"); tmov.heading("v", text="VARIAÇÃO R$")
        tmov.column("p", width=430, anchor="w"); tmov.column("v", width=130, anchor="center")
        tmov.tag_configure("UP", background="#DFF5DF"); tmov.tag_configure("DN", background="#FFE3E3")
        tmov.grid(row=3, column=0, columnspan=2, sticky="nsew")
        ttk.Label(corpo, text="Top 5 mais vendidos de cada mês", font=("Segoe UI", 11, "bold")).grid(
            row=4, column=0, columnspan=2, sticky="w", pady=(6, 0))
        ttop = ttk.Treeview(corpo, columns=("p", "a", "b"), show="headings", height=5)
        ttop.heading("p", text="PRODUTO"); ttop.heading("a", text="MÊS 1 R$"); ttop.heading("b", text="MÊS 2 R$")
        ttop.column("p", width=330, anchor="w"); ttop.column("a", width=120, anchor="center")
        ttop.column("b", width=120, anchor="center")
        ttop.grid(row=5, column=0, columnspan=2, sticky="nsew")
        lbl_novos = tk.Label(corpo, text="", bg=FUNDO, font=("Segoe UI", 10), anchor="w",
                             justify="left", wraplength=640)
        lbl_novos.grid(row=4, column=0, columnspan=2, sticky="w", pady=4)

        def comparar():
            pa_nome, pb_nome = c1.get(), c2.get()
            a, b = h[pa_nome]["produtos"], h[pb_nome]["produtos"]
            ta, tb = h[pa_nome]["total"], h[pb_nome]["total"]
            chip_a.config(text=f"■ {pa_nome[:24]}")
            chip_b.config(text=f"■ {pb_nome[:24]}")
            tv.heading("pa", text=pa_nome[:18]); tv.heading("pb", text=pb_nome[:18])
            for i in tv.get_children():
                tv.delete(i)
            qa = {c: sum(1 for x in a.values() if x["classe"] == c) for c in "ABC"}
            qb = {c: sum(1 for x in b.values() if x["classe"] == c) for c in "ABC"}
            sa = h[pa_nome].get("saude", {}) or {}
            sb = h[pb_nome].get("saude", {}) or {}
            ca, cb = sa.get("cancelados", 0), sb.get("cancelados", 0)
            qqa = sum(x["qtd"] for x in a.values()); qqb = sum(x["qtd"] for x in b.values())
            tka, tkb = (ta/qqa if qqa else 0), (tb/qqb if qqb else 0)
            tem_canc = bool(sa.get("total") and sb.get("total"))
            pts_a = pts_b = 0

            def _lucro_est(prods):
                tot = 0.0
                for x in prods.values():
                    pm = x["fat"]/x["qtd"] if x["qtd"] else 0
                    if pm:
                        tot += x["fat"] - taxa_shopee(pm)*x["qtd"]
                return round(tot, 2)

            la, lb = _lucro_est(a), _lucro_est(b)
            tv.insert("", "end", values=("Faturamento", dinheiro(ta), dinheiro(tb)))
            if la == lb:
                tv.insert("", "end", values=("Lucro estimado*", dinheiro(la), dinheiro(lb)))
            else:
                pts_a, pts_b = (2, 0) if la > lb else (0, 2)
                tv.insert("", "end", values=("Lucro estimado*",
                                             dinheiro(la) + (" ★" if la > lb else ""),
                                             dinheiro(lb) + (" ★" if lb > la else "")))
            for met, va, vb, wa in [
                ("Qtd curva A", qa["A"], qb["A"], qa["A"] >= qb["A"]),
                ("Qtd curva C (menos é melhor)", qa["C"], qb["C"], qa["C"] <= qb["C"]),
                ("Cancelados (menos é melhor)", ca if tem_canc else "—", cb if tem_canc else "—",
                 (ca <= cb) if tem_canc else None),
                ("Ticket médio", dinheiro(tka), dinheiro(tkb), tka >= tkb),
            ]:
                if wa is None:
                    tv.insert("", "end", values=(met, va, vb))
                elif wa:
                    pts_a += 1
                    tv.insert("", "end", values=(met, f"{va} ★", vb))
                else:
                    pts_b += 1
                    tv.insert("", "end", values=(met, va, f"{vb} ★"))
            if pts_a == pts_b:
                banner.config(text="EMPATE TÉCNICO nos critérios.")
            else:
                w = pa_nome if pts_a > pts_b else pb_nome
                banner.config(text=f"MELHOR MÊS: {w}  ({max(pts_a, pts_b)} x {min(pts_a, pts_b)} nos critérios)")
            topa = max(a.items(), key=lambda kv: kv[1]["fat"])[0] if a else "—"
            topb = max(b.items(), key=lambda kv: kv[1]["fat"])[0] if b else "—"
            tv.insert("", "end", values=("Campeão", topa[:28], topb[:28]))
            var = (tb-ta)/ta*100 if ta else 0
            if topa == topb:
                det.config(text=f"Faturamento {var:+.1f}% · {topa[:55]} continua campeão. Proteja com estoque e ads.")
            else:
                det.config(text=f"Faturamento {var:+.1f}% · Antes: {topa[:40]}  →  Agora: {topb[:40]} (virou o jogo).")
            for _t in (tsub, tcai, tmov):
                for _i in _t.get_children():
                    _t.delete(_i)
            subiram = sorted([p for p in b if p in a and a[p]["classe"] == "C" and b[p]["classe"] in "AB"])
            caíram = sorted([p for p in b if p in a and a[p]["classe"] == "A" and b[p]["classe"] in "BC"])
            for p in subiram[:20]:
                tsub.insert("", "end", values=(p[:48], f"C → {b[p]['classe']}"), tags=("UP",))
            if not subiram:
                tsub.insert("", "end", values=("Ninguém subiu de curva", "—"))
            for p in caíram[:20]:
                tcai.insert("", "end", values=(p[:48], f"{a[p]['classe']} → {b[p]['classe']}"), tags=("DN",))
            if not caíram:
                tcai.insert("", "end", values=("Ninguém caiu de curva", "—"))
            novos_c = sorted([p for p in b if (p not in a or a[p]["classe"] in "AB") and b[p]["classe"] == "C"])
            if novos_c:
                lbl_novos.config(text=f"Novos na C ({len(novos_c)}): " + ", ".join(p[:26] for p in novos_c[:6]))
            else:
                lbl_novos.config(text="")
            mov = sorted(((b[p]["fat"]-a[p]["fat"], p) for p in b if p in a), reverse=True)
            for v, p in mov[:3]:
                tmov.insert("", "end", values=(p[:55], f"{v:+.0f}".replace(",", ".")), tags=("UP",))
            for v, p in mov[-3:][::-1]:
                tmov.insert("", "end", values=(p[:55], f"{v:+.0f}".replace(",", ".")), tags=("DN",))
            top5a = sorted(a.items(), key=lambda kv: kv[1]["fat"], reverse=True)[:5]
            top5b = sorted(b.items(), key=lambda kv: kv[1]["fat"], reverse=True)[:5]
            nomes = []
            for p, _v in top5a + top5b:
                if p not in nomes:
                    nomes.append(p)
            fa5 = {p: v["fat"] for p, v in top5a}
            fb5 = {p: v["fat"] for p, v in top5b}
            for _i in ttop.get_children():
                ttop.delete(_i)
            for p in nomes[:10]:
                ttop.insert("", "end", values=(p[:55],
                    f"{fa5.get(p, 0):,.0f}".replace(",", ".") if p in fa5 else "—",
                    f"{fb5.get(p, 0):,.0f}".replace(",", ".") if p in fb5 else "—"))
        c1.bind("<<ComboboxSelected>>", lambda e: comparar())
        c2.bind("<<ComboboxSelected>>", lambda e: comparar())
        tk.Label(pop, text="Troque os meses acima que tudo atualiza sozinho.", bg=FUNDO,
                 font=("Segoe UI", 9)).pack(anchor="w", padx=10)
        ttk.Button(pop, text="Atualizar", style="Acao.TButton", command=comparar).pack(fill="x", padx=10, pady=(0, 10))
        comparar()

    # ----- DRE real (taxas do relatório) -----
    def dre_popup(self):
        if not self.resultado:
            messagebox.showwarning("Calma lá", "Faça a análise (passo 2) primeiro."); return
        tem_real = any(x.get("taxas", 0) for x in self.resultado)
        pop = tk.Toplevel(self)
        pop.title("DRE real por produto")
        pop.geometry("760x640")
        tk.Label(pop, text="DRE REAL — com as taxas que você PAGOU", bg="#141414", fg="white",
                 font=("Segoe UI", 13, "bold")).pack(fill="x", padx=10, pady=(10, 0), ipady=8)
        tk.Label(pop, text=self._periodo_txt() + ("  |  * = taxa real do relatório" if tem_real else "  |  estimativa oficial (importe o Order para o exato)"),
                 bg=FUNDO, font=("Segoe UI", 9)).pack(fill="x", padx=10)
        tv = self.tabela(pop, [("produto", "PRODUTO", 300, True), ("fat", "FATURA", 90, False),
                               ("tax", "TAXAS", 90, False), ("pct", "TAXA%", 70, False),
                               ("lucro", "LUCRO R$", 90, False), ("mg", "MARGEM", 80, False)], 12)
        tv.tag_configure("RUIM", background="#FFCDD2")
        tv.tag_configure("ATEN", background="#FFF3D6")
        tot_f = tot_t = tot_l = 0.0
        top = sorted(self.resultado, key=lambda d: d["fat"], reverse=True)[:20]
        vereditos = []
        for x in top:
            cu = x.get("custo", 0) or 0
            tx = x["taxas"] if x.get("taxas", 0) else (taxa_shopee(x["fat"]/x["qtd"], False, 0)*x["qtd"] if x["qtd"] else 0)
            lucro = x["fat"] - tx - cu*x["qtd"]
            mg = lucro/x["fat"]*100 if x["fat"] else 0
            take = tx/x["fat"]*100 if x["fat"] else 0
            tot_f += x["fat"]; tot_t += tx; tot_l += lucro
            tag = ("RUIM",) if (cu and mg < 8) else (("ATEN",) if take > 32 else ())
            tv.insert("", "end", values=(x["produto"][:48], dinheiro(x["fat"]), dinheiro(tx),
                                         f"{take:.1f}%", dinheiro(lucro), f"{mg:.0f}%" if cu else "s/ custo"),
                      tags=tag)
            if cu and mg < 0:
                ideal = preco_ideal_shopee(cu, 0.10, 0.22, imp=0.04, dev=0.03)
                vereditos.append(f"PREJUÍZO: {x['produto'][:45]} perde {dinheiro(-lucro)}. Suba para {dinheiro(ideal)} ou pause.")
            elif cu and mg < 8:
                vereditos.append(f"NO LIMITE: {x['produto'][:45]} com {mg:.0f}% — suba o preço ou corte custo.")
        box = tk.Frame(pop, bg=ROSA_CLARO)
        box.pack(fill="x", padx=10, pady=8)
        tk.Label(box, text="CONSIDERAÇÃO FINAL — o que corrigir e como", bg=ROSA_CLARO, fg=ROSA_ESCURO,
                 font=("Segoe UI", 12, "bold")).pack(anchor="w", padx=10, pady=(8, 2))
        take_geral = tot_t/tot_f*100 if tot_f else 0
        linhas = [f"Top 20: fatura {dinheiro(tot_f)} · taxas {dinheiro(tot_t)} ({take_geral:.1f}%) · lucro {dinheiro(tot_l)}."]
        linhas += vereditos[:3]
        if take_geral > 30:
            linhas.append("Take rate acima de 30%: revise preços e compare Shopee x ML no botão da Precificação.")
        if not vereditos and take_geral <= 30:
            linhas.append("DRE saudável. Foque a verba nos campeões (classe A).")
        if not tem_real:
            linhas.append("Importe o relatório Order para taxas exatas por produto.")
        for ln in linhas:
            tk.Label(box, text="•  " + ln, bg=ROSA_CLARO, fg=CINZA_TXT, wraplength=700,
                     justify="left", font=("Segoe UI", 11)).pack(anchor="w", padx=10, pady=1)
        tk.Label(box, text="", bg=ROSA_CLARO).pack(pady=2)

    # ----- comparador Shopee x ML -----
    def comp_popup(self):
        if not self.resultado:
            messagebox.showwarning("Calma lá", "Faça a análise (passo 2) primeiro."); return
        pop = tk.Toplevel(self)
        pop.title("Shopee x Mercado Livre")
        pop.geometry("900x560")
        imp, dev = self.p_imposto()
        tv = self.tabela(pop, [("produto", "PRODUTO", 280, True), ("preco", "PREÇO", 85, False),
                               ("sh", "SH LÍQUIDO", 95, False), ("shm", "SH MARGEM", 85, False),
                               ("ml", "ML LÍQUIDO", 95, False), ("mlm", "ML MARGEM", 85, False),
                               ("melhor", "MELHOR", 90, False)], 14)
        for x in sorted(self.resultado, key=lambda d: d["fat"], reverse=True)[:15]:
            pm = x["fat"]/x["qtd"] if x["qtd"] else 0
            cu = x.get("custo", 0) or 0
            tsh = taxa_shopee(pm)
            tml = taxa_ml(pm, premium=False)
            lsh, lml = pm-tsh, pm-tml
            msh = ((lsh-cu)/pm - imp - dev) if pm else 0
            mml = ((lml-cu)/pm - imp - dev) if pm else 0
            if not cu:
                melhor = "—"
            elif abs(msh-mml) < 0.02:
                melhor = "Empate"
            else:
                melhor = "Shopee" if msh > mml else "ML"
            tv.insert("", "end", values=(x["produto"][:45], dinheiro(pm), dinheiro(lsh),
                f"{msh*100:.0f}%" if cu else "s/ custo", dinheiro(lml),
                f"{mml*100:.0f}%" if cu else "s/ custo", melhor))
        tk.Label(pop, text="ML calculado no Clássico, frete leve. Informe custos no passo 3 para margem exata.",
                 bg=FUNDO, font=("Segoe UI", 9)).pack(anchor="w", padx=10, pady=(0, 8))

    # ----- calendário de campanhas -----
    def camp_popup(self):
        import datetime as _dt
        import calendar as _cal
        hoje = _dt.date.today()

        def segundo_domingo(ano, mes):
            c = _cal.monthcalendar(ano, mes)
            doms = [s[6] for s in c if s[6]]
            return _dt.date(ano, mes, doms[1])

        def black_friday(ano):
            c = _cal.monthcalendar(ano, 11)
            sexs = [s[4] for s in c if s[4]]
            return _dt.date(ano, 11, sexs[3])

        fixas = [(9, 9, "9.9 Super Shopping Day"), (10, 10, "10.10"), (11, 11, "11.11"),
                 (12, 12, "12.12"), (12, 25, "Natal"), (1, 1, "Ano Novo"),
                 (3, 15, "Dia do Consumidor"), (10, 12, "Dia das Crianças")]
        eventos = []
        for ano in (hoje.year, hoje.year+1):
            for m, d, nome in fixas:
                eventos.append((_dt.date(ano, m, d), nome))
            eventos.append((segundo_domingo(ano, 5), "Dia das Mães"))
            eventos.append((segundo_domingo(ano, 8), "Dia dos Pais"))
            eventos.append((black_friday(ano), "Black Friday"))
        eventos = sorted((dt, n) for dt, n in eventos if dt >= hoje)[:6]
        L = ["CALENDÁRIO DE CAMPANHAS — o que preparar", "="*60]
        for dt, nome in eventos:
            dias = (dt-hoje).days
            if dias > 30:
                chk = "definir verba e escolher os campeões da campanha."
            elif dias > 15:
                chk = "subir estoque, fotos e títulos dos escolhidos."
            elif dias > 7:
                chk = "ativar cupons, kits e revisar preços."
            elif dias > 0:
                chk = "última chamada: estoque, preço e ads ligados."
            else:
                chk = "É HOJE: monitore estoque e verba de hora em hora."
            L.append(f"\n{dt.strftime('%d/%m')} ({dias} dias) — {nome}\n  Agora: {chk}")
        L.append("\nVale para Shopee e ML: as datas dobradas (9.9, 11.11...) concentram o tráfego.")
        pop = tk.Toplevel(self)
        pop.title("Calendário de campanhas")
        pop.geometry("620x520")
        out = scrolledtext.ScrolledText(pop, wrap="word", font=("Segoe UI", 12))
        out.pack(fill="both", expand=True, padx=10, pady=10)
        out.insert("end", "\n".join(L))
        out.config(state="disabled")

    # ----- foco C: atenção máxima -----
    def fococ_popup(self):
        if not self.resultado:
            messagebox.showwarning("Calma lá", "Faça a análise (passo 2) primeiro."); return
        cs = [x for x in self.resultado if x.get("classe") == "C"]
        if not cs:
            messagebox.showinfo("Boa notícia", "Você não tem produtos na curva C!"); return
        pop = tk.Toplevel(self)
        pop.title("Foco C — atenção máxima")
        pop.geometry("820x620")
        tv = self.tabela(pop, [("produto", "PRODUTO", 280, True), ("vend", "VENDAS", 70, False),
                               ("est", "ESTOQUE", 75, False), ("margem", "MARGEM", 75, False),
                               ("decisao", "DECISÃO", 130, True), ("acao", "O QUE FAZER", 230, True)], 8)
        tv.tag_configure("LIQ", background="#FFCDD2")
        tv.tag_configure("TST", background="#FFF3D6")
        tv.tag_configure("OBS", background="#F0F0F0")
        tv.bind("<<TreeviewSelect>>", lambda e: self._fococ_detalhe(tv, out))
        out = scrolledtext.ScrolledText(pop, wrap="word", height=6, font=("Segoe UI", 12))
        out.pack(fill="x", padx=10, pady=(0, 8))
        out.tag_configure("N", font=("Segoe UI", 11, "bold"))
        tv.tag_configure("LIQ", background="#FFCDD2")
        tv.tag_configure("TST", background="#FFF3D6")
        tv.tag_configure("OBS", background="#F0F0F0")
        dias = self._dias_periodo()
        n_liq = n_tst = cap_parado = 0.0
        for x in sorted(cs, key=lambda d: d.get("est", 0)*((d.get("custo", 0) or 0) or (d["fat"]/d["qtd"] if d["qtd"] else 0)), reverse=True):
            pm = x["fat"]/x["qtd"] if x["qtd"] else 0
            cu = x.get("custo", 0) or 0
            mg = ((pm-cu)/pm if pm and cu else 0)
            daily = x["qtd"]/dias if dias else 0
            giro = (x.get("est", 0)/daily) if daily else 999
            if (cu and mg < 0.05) or (x.get("est", 0) > 0 and daily > 0 and giro > 120):
                dec, tag = "LIQUIDAR", "LIQ"
                acao = "Queima 20-30% off, não repõe; anuncia só em kit até zerar."
                n_liq += 1
            elif len(x["produto"]) < 40 or not cu:
                dec, tag = "TESTAR", "TST"
                acao = "Arrumar título/fotos, informar custo e testar com pouca verba."
            else:
                dec, tag = "KIT", "OBS"
                acao = "Manter só dentro de kit com campeão."
            cap_parado += x.get("est", 0)*(cu or pm)
            tv.insert("", "end", values=(x["produto"][:45], x["qtd"], x.get("est", 0),
                                         f"{mg*100:.0f}%" if cu else "s/ custo", dec, acao), tags=(tag,))
            if dec == "TESTAR":
                n_tst += 1
        tk.Label(pop, text=f"{len(cs)} produtos C | LIQUIDAR: {n_liq} | TESTAR: {n_tst} | Capital parado em estoque: {dinheiro(cap_parado)}",
                 bg=FUNDO, font=("Segoe UI", 11, "bold")).pack(fill="x", padx=10, pady=(0, 8))

    def _fococ_detalhe(self, tv, out):
        sel = tv.selection()
        if not sel:
            return
        nome = tv.item(sel[0])["values"][0]
        meus = [x for x in self.resultado if x["produto"][:55] == nome or x["produto"] == nome]
        if not meus:
            meus = [x for x in self.resultado if x["produto"][:45] in nome]
        if not meus:
            return
        x = meus[0]
        pm = x["fat"]/x["qtd"] if x["qtd"] else 0
        out.delete("1.0", "end")
        out.insert("end", f"ESTRATÉGIA — {x['produto'][:60]}\n", "N")
        out.insert("end", "O QUE ALTERAR:\n", "N")
        alt = []
        if len(x["produto"]) < 40:
            alt.append("título curto: reescreva com 80+ letras (marca + modelo + medida + uso)")
        alt.append("fotos: complete 7 (capa branca, zoom, medida com régua, em uso, caixa)")
        if x.get("custo"):
            alt.append(f"preço {dinheiro(pm)}: confira margem e compare com a mediana do nicho")
        else:
            alt.append("informe o custo no passo 3 para validar a margem")
        for a in alt:
            out.insert("end", f"• {a}\n")
        out.insert("end", "O QUE ANALISAR (14 dias de teste):\n", "N")
        for a in ["visitas e CTR do anúncio (CTR saudável > 1%)",
                  "conversão visitas→vendas (saudável > 2%)",
                  "posição de preço x 4 mais vendidos (aba 6)",
                  "margem real após taxas + ads"]:
            out.insert("end", f"• {a}\n")
        out.insert("end", "QUANDO EXCLUIR E CRIAR OUTRO:\n", "N")
        for a in ["14 dias com visitas e ZERO vendas após os ajustes acima",
                  "CTR abaixo de 0,5% após 3 mil impressões",
                  "margem negativa sem como subir o preço",
                  "produto parado há 60+ dias: exclua, liquide o estoque e NÃO crie igual — crie como KIT",
                  "ao criar outro: título novo + capa nova (nunca duplique o esgotado)"]:
            out.insert("end", f"• {a}\n")

    # ----- evolução com gráficos + WhatsApp -----
    def evo_popup(self):
        h = self._hist_load()
        if len(h) < 1:
            messagebox.showinfo("Sem histórico",
                "Faça análises com período em meses diferentes.\nCada análise é salva sozinha."); return
        pop = tk.Toplevel(self)
        pop.title("Evolução mês a mês")
        pop.geometry("720x640")
        top = ttk.Frame(pop); top.pack(fill="x", padx=10, pady=8)
        ttk.Label(top, text="Ver:").pack(side="left")
        cb = ttk.Combobox(top, values=["Faturamento R$", "Qtd classe A", "Qtd classe C"], width=18, state="readonly")
        cb.set("Faturamento R$"); cb.pack(side="left", padx=6)
        cv = tk.Canvas(pop, width=680, height=300, bg="white", highlightthickness=1,
                       highlightbackground="#D8CFB8")
        cv.pack(padx=10, pady=4)
        out = scrolledtext.ScrolledText(pop, wrap="word", height=7, font=("Segoe UI", 12))
        out.pack(fill="x", padx=10, pady=6)
        out.tag_configure("N", font=("Segoe UI", 12, "bold"))

        def serie(metrica):
            pers = sorted(h.keys())
            vals = []
            for p in pers:
                pr = h[p]["produtos"]
                if metrica.startswith("Fat"):
                    vals.append(h[p]["total"])
                elif "classe A" in metrica:
                    vals.append(sum(1 for x in pr.values() if x["classe"] == "A"))
                else:
                    vals.append(sum(1 for x in pr.values() if x["classe"] == "C"))
            return pers, vals

        def desenhar(_e=None):
            pers, vals = serie(cb.get())
            cv.delete("all")
            W, H, m = 680, 300, 40
            mx = max(vals) or 1
            n = len(vals)
            larg = min(90, (W-2*m)//max(n, 1) - 14)
            for i, (p, v) in enumerate(zip(pers, vals)):
                x0 = m + i*((W-2*m)//max(n, 1)) + 7
                alt = (H-2*m-24)*(v/mx)
                cor = "#C9A227" if i == n-1 else "#3D0A1A"
                cv.create_rectangle(x0, H-m-alt, x0+larg, H-m, fill=cor, outline="")
                txt = f"{v:,.0f}".replace(",", ".") if "Fat" in cb.get() else str(v)
                cv.create_text(x0+larg//2, H-m-alt-12, text=txt, font=("Segoe UI", 10, "bold"))
                cv.create_text(x0+larg//2, H-m+16, text=p[:10], font=("Segoe UI", 9))
            if n > 1:
                var = (vals[-1]-vals[0])/vals[0]*100 if vals[0] else 0
                cv.create_text(W//2, 20, text=f"{cb.get()}: {var:+.1f}% no período",
                               font=("Segoe UI", 12, "bold"), fill="#3D0A1A")

        def resumo_whats():
            pers, vals = serie("Faturamento R$")
            tot = vals[-1] if vals else 0
            var = f" ({(vals[-1]-vals[0])/vals[0]*100:+.1f}% vs {pers[0][:10]})" if len(vals) > 1 and vals[0] else ""
            ca = sum(1 for x in self.resultado if x.get("classe") == "A") if self.resultado else 0
            txt = (f"*CENTRAL REH* — resumo {pers[-1][:10] if pers else ''}\n"
                   f"Faturamento: R$ {tot:,.2f}{var}\n".replace(",", ".") +
                   f"Campeões (A): {ca} produtos\n"
                   "Detalhes no programa. Bora vender!")
            self.clipboard_clear(); self.clipboard_append(txt); self.update()
            try:
                webbrowser.open("https://web.whatsapp.com/")
            except Exception:
                pass
            messagebox.showinfo("Pronto", "Texto copiado! Cole no WhatsApp que abriu.")

        cb.bind("<<ComboboxSelected>>", desenhar)
        bar = ttk.Frame(pop); bar.pack(fill="x", padx=10, pady=(0, 10))
        ttk.Button(bar, text="Copiar resumo p/ WhatsApp", command=resumo_whats).pack(side="left", padx=2)
        desenhar()
        out.delete("1.0", "end")
        out.insert("end", "COMO LER:\n", "N")
        out.insert("end", "Barra dourada = mês atual. Subindo = crescendo.\nTroque a métrica acima para ver campeões (A) e parados (C) ao longo dos meses.")

    # ----- envio ao Full -----
    def full_popup(self):
        if not self.resultado:
            messagebox.showwarning("Calma lá", "Faça a análise (passo 2) primeiro."); return
        dias = self._dias_periodo()
        nome_camp, dias_camp, extra = proxima_campanha()
        pop = tk.Toplevel(self)
        pop.title("Enviar para o Full")
        pop.geometry("820x620")
        tk.Label(pop, text=f"Campanha chegando: {nome_camp} em {dias_camp} dias → reserva extra de {extra*100:.0f}% no cálculo. Envie com 7 dias de antecedência.",
                 bg=ROSA_CLARO, fg=CINZA_TXT, font=("Segoe UI", 11, "bold"), wraplength=780,
                 justify="left").pack(fill="x", padx=10, pady=(10, 4))
        v15, v30 = tk.BooleanVar(value=True), tk.BooleanVar(value=True)
        fop = ttk.Frame(pop); fop.pack(fill="x", padx=10, pady=2)
        ttk.Label(fop, text="Atender por:").pack(side="left")
        ttk.Checkbutton(fop, text="15 dias", variable=v15, command=lambda: redesenhar()).pack(side="left", padx=4)
        ttk.Checkbutton(fop, text="30 dias", variable=v30, command=lambda: redesenhar()).pack(side="left", padx=2)
        box = ttk.Frame(pop); box.pack(fill="both", expand=True, padx=10, pady=4)
        lbl_tot = tk.Label(pop, text="", bg=FUNDO, font=("Segoe UI", 11, "bold"))
        lbl_tot.pack(fill="x", padx=10, pady=(0, 8))
        base_linhas = []
        for x in [r for r in self.resultado if r.get("classe") == "A"]:
            daily = x["qtd"]/dias if dias else 0
            if daily <= 0:
                continue
            est = x.get("est", 0) or 0
            base_linhas.append((est/daily, x["produto"], daily, est))

        def redesenhar():
            for w in box.winfo_children():
                w.destroy()
            ds = ([15] if v15.get() else []) + ([30] if v30.get() else [])
            if not ds:
                lbl_tot.config(text="Marque 15 e/ou 30 dias.")
                return
            cols = [("produto", "PRODUTO (curva A)", 250, True), ("vd", "VENDA/DIA", 70, False),
                    ("est", "ESTOQUE", 65, False), ("cob", "COBERTURA", 85, False)]
            for d in ds:
                cols += [(f"e{d}", f"ENVIAR {d}D", 80, False), (f"a{d}", f"ATENDE {d}D", 80, False)]
            tv = ttk.Treeview(box, columns=[c[0] for c in cols], show="headings", height=13)
            for cid, tit, larg, esq in cols:
                tv.heading(cid, text=tit)
                tv.column(cid, width=larg, minwidth=larg//2, stretch=True, anchor="w" if esq else "center")
            tv.tag_configure("URG", background="#FFCDD2")
            tv.tag_configure("OKK", background="#DFF5DF")
            sb = ttk.Scrollbar(box, orient="vertical", command=tv.yview)
            tv.configure(yscrollcommand=sb.set)
            tv.pack(side="left", fill="both", expand=True)
            sb.pack(side="right", fill="y")
            tot = {d: 0 for d in ds}
            linhas = []
            for cob, prod, daily, est in base_linhas:
                qs = {}
                for d in ds:
                    q = max(0, int(d*daily*(1+extra) - est + 0.5))
                    qs[d] = q
                    tot[d] += q
                linhas.append((max(qs.values()), cob, prod, daily, est, qs))
            linhas.sort(reverse=True)
            for _mx, cob, prod, daily, est, qs in linhas:
                tag = ("URG",) if cob < 14 else ()
                vals = [prod[:50], f"{daily:.1f}", est, f"{cob:.0f} dias"]
                for d in ds:
                    vals += [qs[d] if qs[d] else "ok", f"{(est+qs[d])/daily:.0f} dias"]
                tv.insert("", "end", values=vals, tags=tag)
            lbl_tot.config(text=f"{len(linhas)} campeões | " +
                                "  ".join(f"Total {d}d: {tot[d]} un" for d in ds) +
                                f" (+ campanha {nome_camp})")

        redesenhar()

    # ----- montador de kits -----
    def kit_popup(self):
        if not self.resultado:
            messagebox.showwarning("Calma lá", "Faça a análise (passo 2) primeiro."); return
        pop = tk.Toplevel(self)
        pop.title("Montar kit A+C")
        pop.geometry("660x600")
        la = [x["produto"] for x in self.resultado if x["classe"] == "A"][:15]
        lc = [x["produto"] for x in self.resultado if x["classe"] == "C"]
        f = ttk.Frame(pop); f.pack(fill="x", padx=10, pady=6)
        f.columnconfigure(1, weight=1)
        ttk.Label(f, text="Campeão (A):").grid(row=0, column=0, sticky="w", pady=3)
        ca = ttk.Combobox(f, values=la, width=45, state="readonly")
        if la: ca.set(la[0])
        ca.grid(row=0, column=1, sticky="we", padx=6, pady=3)
        ttk.Label(f, text="Parado (C):").grid(row=1, column=0, sticky="w", pady=3)
        cc = ttk.Combobox(f, values=lc, width=45, state="readonly")
        if lc: cc.set(lc[0])
        cc.grid(row=1, column=1, sticky="we", padx=6, pady=3)
        g = ttk.Frame(pop); g.pack(fill="x", padx=10, pady=4)
        ttk.Label(g, text="Desconto %:").pack(side="left")
        e_desc = ttk.Entry(g, width=8); e_desc.insert(0, "12"); e_desc.pack(side="left", padx=4)
        ttk.Label(g, text="Plataforma:").pack(side="left", padx=(8, 0))
        c_plat = ttk.Combobox(g, values=["Shopee", "ML Clássico", "ML Premium"], width=12, state="readonly")
        c_plat.set("Shopee"); c_plat.pack(side="left", padx=4)
        cf = ttk.Frame(pop); cf.pack(fill="x", padx=10, pady=2)
        ttk.Label(cf, text="Custo A R$:").pack(side="left")
        e_caua = ttk.Entry(cf, width=10); e_caua.pack(side="left", padx=4)
        ttk.Label(cf, text="Custo C R$:").pack(side="left")
        e_cauc = ttk.Entry(cf, width=10); e_cauc.pack(side="left", padx=4)
        tk.Label(cf, text="(já vem preenchido; ajuste se precisar)", bg=FUNDO, font=("Segoe UI", 9)).pack(side="left", padx=4)
        out = scrolledtext.ScrolledText(pop, wrap="word", font=("Segoe UI", 12))
        out.pack(fill="both", expand=True, padx=10, pady=6)
        out.tag_configure("N", font=("Segoe UI", 12, "bold"))

        def dados(nome):
            for x in self.dados:
                if x["produto"] == nome:
                    return x
            return None

        def simular():
            A, C = dados(ca.get()), dados(cc.get())
            if not A or not C:
                out.delete("1.0", "end"); out.insert("end", "Escolha os dois produtos."); return
            try: desc = num(e_desc.get())/100
            except ValueError: desc = 0.12
            pa = A["fat"]/A["qtd"] if A["qtd"] else 0
            pc = C["fat"]/C["qtd"] if C["qtd"] else 0
            soma, kit = pa+pc, round((pa+pc)*(1-desc), 2)
            try:
                cau = num(e_caua.get()) if e_caua.get().strip() else (A.get("custo", 0) or 0)
                ccu = num(e_cauc.get()) if e_cauc.get().strip() else (C.get("custo", 0) or 0)
            except ValueError:
                cau, ccu = A.get("custo", 0) or 0, C.get("custo", 0) or 0
            for _base, _nome, _cu in ((A, A["produto"], cau), (C, C["produto"], ccu)):
                for _x in self.dados + self.resultado:
                    if _x["produto"] == _nome:
                        _x["custo"] = _cu
                _reg = self.loja.get(_nome, {})
                _reg["custo"] = _cu
                self.loja[_nome] = _reg
            self._loja_save()
            plat = c_plat.get()
            fee = taxa_shopee(kit) if plat == "Shopee" else taxa_ml(kit, "Premium" in plat)
            imp, dev = self.p_imposto()
            margem = ((kit-cau-ccu-fee)/kit - imp - dev) if kit else 0
            tit = " ".join(f"Kit {A['produto']} + {C['produto']}".split())
            if len(tit) > 120:
                tit = " ".join(f"Kit {A['produto'][:55]} + {C['produto'][:40]}".split())[:120]
            veredito = ("ÓTIMO — pode anunciar o kit." if margem >= 0.15 else
                        ("OK — teste com pouca verba." if margem >= 0.08 else
                         "RUIM — diminua o desconto ou troque o par."))
            out.delete("1.0", "end")
            out.insert("end", "TÍTULO SUGERIDO DO KIT:\n", "N")
            out.insert("end", f"{tit}\n({len(tit)} letras, base nos nomes que já vendem)\n\n")
            out.insert("end", "PREÇO DE VENDA SUGERIDO:\n", "N")
            out.insert("end", f"De {dinheiro(soma)} por apenas {dinheiro(kit)} ({desc*100:.0f}% off nos preços já vendidos)\n\n")
            out.insert("end", f"Taxas {plat}: {dinheiro(fee)} | Custos: {dinheiro(cau+ccu)}\n")
            out.insert("end", f"Lucro do kit: {dinheiro(kit-cau-ccu-fee)} ({margem*100:.1f}% já com imposto e perdas)\n\n")
            out.insert("end", f"VEREDITO: {veredito}\n", "N")

        ttk.Button(pop, text="Simular kit", style="Acao.TButton", command=simular).pack(fill="x", padx=10, pady=(0, 4))
        def recarregar():
            A, C = dados(ca.get()), dados(cc.get())
            e_caua.delete(0, "end")
            e_cauc.delete(0, "end")
            if A:
                e_caua.insert(0, f"{(A.get('custo', 0) or 0):.2f}".replace(".", ","))
            if C:
                e_cauc.insert(0, f"{(C.get('custo', 0) or 0):.2f}".replace(".", ","))
            simular()
        ca.bind("<<ComboboxSelected>>", lambda e: recarregar())
        cc.bind("<<ComboboxSelected>>", lambda e: recarregar())
        c_plat.bind("<<ComboboxSelected>>", lambda e: simular())
        e_desc.bind("<KeyRelease>", lambda e: simular())
        sug = ["SUGESTÕES AUTOMÁTICAS (campeão + encalhado):"]
        enc = sorted([x for x in self.resultado if x["classe"] == "C"], key=lambda x: x.get("est", 0), reverse=True)[:3]
        for i, c in enumerate(enc):
            a = next((x for x in self.resultado if x["classe"] == "A"), None)
            if a:
                sug.append(f"• {a['produto'][:40]} + {c['produto'][:40]}")
        out.insert("end", "\n".join(sug))
        if la and lc:
            recarregar()

    # ----- cancelados -----
    def _tela_canc(self):
        t = self.tX
        self.dica(t, "Quem está cancelando, por quê e o que fazer. Precisa do relatório de PEDIDOS (Order...xlsx).",
                  "curva")
        self.x_info = tk.Label(t, text="Importe o relatório de pedidos para ver os cancelados.",
                               bg=FUNDO, font=("Segoe UI", 12, "bold"))
        self.x_info.pack(anchor="w", padx=14, pady=4)
        self.xtree = self.tabela(t, [("produto", "PRODUTO", 280, True), ("cl", "CL", 45, False),
                                     ("qtd", "CANCEL.", 80, False), ("valor", "R$ PERDIDO", 100, False),
                                     ("motivo", "MOTIVO PRINCIPAL", 220, True),
                                     ("sol", "COMO RESOLVER", 250, True)], 11)
        self.xtree.tag_configure("A", background="#DFF5DF"); self.xtree.tag_configure("B", background="#FFF3D6")
        self.xtree.tag_configure("C", background="#F0F0F0")

    def _canc_fill(self):
        s = getattr(self, "saude", None) or {}
        for i in self.xtree.get_children():
            self.xtree.delete(i)
        if not s or not s.get("por_produto"):
            self.x_info.config(text="Importe o relatório de pedidos para ver os cancelados.")
            return
        pp = s["por_produto"]
        classes = {x["produto"]: x.get("classe", "") for x in self.resultado}
        tot_c = sum(v["qtd"] for v in pp.values())
        linhas = []
        for prod, v in pp.items():
            mot = max(v["motivos"].items(), key=lambda kv: kv[1])[0] if v["motivos"] else "—"
            linhas.append((v["qtd"], prod, classes.get(prod, "—"), v["valor"], mot))
        linhas.sort(reverse=True)
        resumo = {"A": [0, 0.0], "B": [0, 0.0], "C": [0, 0.0], "—": [0, 0.0]}
        for _q, prod, cl, valor, _m in linhas:
            if cl in resumo:
                resumo[cl][0] += _q; resumo[cl][1] += valor
        self.x_info.config(
            text=f"{s['cancelados']} cancelados ({dinheiro(s['v_cancelado'])} perdidos)  |  "
                 f"A: {resumo['A'][0]} ({dinheiro(resumo['A'][1])})   "
                 f"B: {resumo['B'][0]} ({dinheiro(resumo['B'][1])})   "
                 f"C: {resumo['C'][0]} ({dinheiro(resumo['C'][1])})")
        top = linhas[0]
        self.x_info.config(text=self.x_info.cget("text") +
                                 f"  |  Mais cancelado: {top[1][:40]} ({top[0]}x)")
        for q, prod, cl, valor, mot in linhas[:30]:
            tag = (cl,) if cl in "ABC" else ()
            self.xtree.insert("", "end", values=(prod[:55], cl, q, dinheiro(valor),
                                                 mot[:60], solucao_motivo(mot)[:70]), tags=tag)

    # ----- 3 planilhas: comparar meses/plataformas -----
    def _ler_arquivo(self, arq):
        """Lê 1 arquivo sem mexer na tela. Retorna (dados, periodo, saude)."""
        import re as _re, datetime as _dt
        ext = arq.lower().rsplit(".", 1)[-1]
        per = ""
        m = _re.search(r"(\d{8})_(\d{8})", arq)
        if m:
            try:
                per = (_dt.datetime.strptime(m.group(1), "%Y%m%d").strftime("%d/%m/%Y") + " a " +
                       _dt.datetime.strptime(m.group(2), "%Y%m%d").strftime("%d/%m/%Y"))
            except ValueError:
                per = ""
        if ext == "xlsx":
            dados = ler_excel(arq)
        elif ext == "csv":
            dados = ler_csv(arq)
        else:
            raise ValueError("Só aceito .xlsx ou .csv.")
        saude = dict(ULTIMOS_PEDIDOS) if ULTIMOS_PEDIDOS else {}
        return dados, per, saude

    def _slot_estilos(self):
        for n in (1, 2, 3):
            if n == self.slot_ativo:
                self.slot_btns[n].config(style="Ativo.TButton")
            elif self.slots[n]:
                self.slot_btns[n].config(style="Acao.TButton")
            else:
                self.slot_btns[n].config(style="TButton")

    def slot_clicar(self, n):
        if self.slots[n] is None:
            arq = filedialog.askopenfilename(filetypes=[("Excel ou CSV", "*.xlsx *.csv")])
            if not arq: return
            try:
                dados, per, saude = self._ler_arquivo(arq)
            except Exception as e:
                messagebox.showerror("Não consegui ler", str(e)); return
            res, tot = classificar(dados)
            nome = arq.rsplit("/", 1)[-1].rsplit("\\", 1)[-1][:30]
            self.slots[n] = {"arq": nome, "dados": dados, "resultado": res, "total": tot,
                             "periodo": per or "sem período", "saude": saude}
            self.slot_lbls[n].config(text=f"{nome[:22]} ({per or 's/ período'})")
            self.slot_btns[n].config(text=f"Ver plan. {n}")
        s = self.slots[n]
        self.dados = s["dados"]
        self.resultado = s["resultado"]
        self.total = s["total"]
        self.saude = s["saude"]
        self.slot_ativo = n
        self.arq_atual = s["arq"]
        self._slot_estilos()
        d1, d2 = (s["periodo"].split(" a ") + ["", ""])[:2]
        self.p_de.delete(0, "end"); self.p_de.insert(0, d1)
        self.p_ate.delete(0, "end"); self.p_ate.insert(0, d2)
        self._refresh()
        self.analisar()
        self.nb.select(self.tA)

    def slots_limpar(self):
        self.slots = {1: None, 2: None, 3: None}
        self.slot_ativo = None
        for n in (1, 2, 3):
            self.slot_btns[n].config(text=f"Planilha {n}")
            self.slot_lbls[n].config(text="vazia")
        self._slot_estilos()

    def slots_comparar(self):
        cheios = [(n, s) for n, s in self.slots.items() if s]
        if len(cheios) < 2:
            messagebox.showwarning("Faltam planilhas",
                "Carregue ao menos 2 planilhas (botões Planilha 1, 2 e 3)."); return
        pop = tk.Toplevel(self)
        pop.title("Comparar planilhas")
        pop.geometry("760x600")
        banner = tk.Label(pop, text="", bg="#DFF5DF", fg="#1E7E34", font=("Segoe UI", 13, "bold"),
                          padx=10, pady=8, wraplength=720, justify="left")
        banner.pack(fill="x", padx=10, pady=(10, 4))
        cols = ["m"] + [f"p{n}" for n, _s in cheios]
        tv = ttk.Treeview(pop, columns=cols, show="headings", height=8)
        tv.heading("m", text="MEDIDA")
        tv.column("m", width=200, anchor="w")
        for n, s in cheios:
            tv.heading(f"p{n}", text=f"Planilha {n}")
            tv.column(f"p{n}", width=150, anchor="center")
        tv.pack(fill="x", padx=10, pady=4)
        out = scrolledtext.ScrolledText(pop, wrap="word", height=6, font=("Segoe UI", 12))
        out.pack(fill="x", padx=10, pady=(0, 10))
        out.tag_configure("N", font=("Segoe UI", 12, "bold"))
        mets, pts = [], {n: 0 for n, _s in cheios}
        for n, s in cheios:
            fat = s["total"]
            qq = sum(x["qtd"] for x in s["dados"]) or 1
            canc = s["saude"].get("cancelados", 0)
            totc = s["saude"].get("total", 0) or 1
            camp = max(s["resultado"], key=lambda x: x["fat"])["produto"][:38] if s["resultado"] else "—"
            mets.append({"n": n, "per": s["periodo"], "fat": fat, "nprod": len(s["dados"]),
                         "ticket": fat/qq, "canc": canc, "txc": canc/totc*100,
                         "taxas": sum(x.get("taxas", 0) for x in s["dados"]), "camp": camp})
        tv.insert("", "end", values=["Período"] + [m["per"][:20] for m in mets])
        tv.insert("", "end", values=["Faturamento"] + [dinheiro(m["fat"]) for m in mets])
        tv.insert("", "end", values=["Produtos"] + [m["nprod"] for m in mets])
        tv.insert("", "end", values=["Ticket"] + [dinheiro(m["ticket"]) for m in mets])
        tv.insert("", "end", values=["Cancelados"] + [f"{m['canc']} ({m['txc']:.1f}%)" for m in mets])
        tv.insert("", "end", values=["Taxas pagas"] + [dinheiro(m["taxas"]) for m in mets])
        tv.insert("", "end", values=["Campeão"] + [m["camp"][:28] for m in mets])
        def ganha(vals, maior=True):
            top = max(vals) if maior else min(vals)
            return [i for i, v in enumerate(vals) if v == top]
        for idx in ([m["fat"] for m in mets], [m["ticket"] for m in mets]):
            for i in ganha(idx, True):
                pts[mets[i]["n"]] += 1
        for idx in ([m["canc"] for m in mets], [m["taxas"] for m in mets]):
            for i in ganha(idx, False):
                pts[mets[i]["n"]] += 1
        win = max(pts, key=lambda n: pts[n])
        todos_ig = len(set(pts.values())) == 1
        banner.config(text="EMPATE TÉCNICO entre as planilhas." if todos_ig else
                      f"MELHOR: PLANILHA {win} ({mets[[m['n'] for m in mets].index(win)]['per']}) — "
                      f"venceu em {pts[win]} de 4 critérios (faturamento, ticket, menos cancelados, menos taxas).")
        out.delete("1.0", "end")
        out.insert("end", "DETALHE POR PLANILHA:\n", "N")
        for m in mets:
            out.insert("end", f"\nPlanilha {m['n']} ({m['per']}): "
                              f"fatura {dinheiro(m['fat'])} com {m['nprod']} produtos; "
                              f"cancela {m['canc']} ({m['txc']:.1f}%); taxas {dinheiro(m['taxas'])}; "
                              f"campeão: {m['camp'][:50]}.\n")

    # ---------- ações ----------
    def _refresh(self, com_classe=False):
        for i in self.tree.get_children(): self.tree.delete(i)
        base = self.resultado if com_classe and self.resultado else self.dados
        flt = getattr(self, "filtro_val", "Todas")
        if com_classe and flt in "ABC":
            base = [x for x in base if x.get("classe") == flt]
        for x in base:
            cl = x.get("classe") or ""
            tag = (cl,) if cl in "ABC" else ()
            self.tree.insert("", "end", values=(x["produto"][:80], x["qtd"], dinheiro(x["fat"]), cl or "—"), tags=tag)
        tot = sum(x["fat"] for x in base)
        a = sum(1 for x in self.resultado if x.get("classe") == "A")
        b = sum(1 for x in self.resultado if x.get("classe") == "B")
        c = sum(1 for x in self.resultado if x.get("classe") == "C")
        detalhe = f"  |  A: {a}  B: {b}  C: {c}" if self.resultado else ""
        onde = f"[{getattr(self, 'arq_atual', '')}] " if getattr(self, "arq_atual", "") else ""
        self.info.config(text=f"{onde}{len(base)} produtos | Somando {dinheiro(tot)}{detalhe}" if base else "Nenhum produto ainda.")
        self._status()

    def imp(self):
        arqs = filedialog.askopenfilenames(filetypes=[("Excel ou CSV", "*.xlsx *.csv"), ("Excel", "*.xlsx"), ("CSV", "*.csv")])
        if not arqs: return
        if isinstance(arqs, str):
            arqs = (arqs,)
        if len(arqs) == 1:
            self._imp_um(arqs[0], avisar=True)
            return
        feitos, falhas = [], []
        for arq in arqs:
            ok, info = self._imp_um(arq, avisar=False)
            (feitos if ok else falhas).append(info)
        self.nb.select(self.tA)
        msg = "Análise mês a mês pronta!\n\n" + "\n".join(f"• {x}" for x in feitos)
        if falhas:
            msg += "\n\nNão entraram:\n" + "\n".join(f"• {x}" for x in falhas)
        msg += "\n\nUse 'Comparar meses' e 'Evolução' na aba 2."
        messagebox.showinfo("Vários meses", msg)
        try:
            self.hist_popup()
        except Exception:
            pass

    def _imp_um(self, arq, avisar=True, forcar_periodo=False):
        """Importa 1 arquivo e já analisa. Retorna (ok, resumo)."""
        import re as _re, datetime as _dt
        nome = arq.rsplit("/", 1)[-1].rsplit("\\", 1)[-1]
        try:
            ext = arq.lower().rsplit(".", 1)[-1]
            if ext == "xlsx":
                try:
                    self.dados = ler_excel(arq, _data_br(self.p_de.get()), _data_br(self.p_ate.get()))
                except ValueError as ve:
                    if "vendas válidas" in str(ve) and (self.p_de.get().strip() or self.p_ate.get().strip()):
                        self.dados = ler_excel(arq)
                        self.p_de.delete(0, "end"); self.p_ate.delete(0, "end")
                        if avisar:
                            messagebox.showwarning("Período diferente",
                                "As datas antigas não batiam com esse arquivo, então usei tudo.\n"
                                "Confira o período detectado abaixo.")
                    else:
                        raise
            elif ext == "csv":
                self.dados = ler_csv(arq)
            else:
                if avisar:
                    messagebox.showwarning("Formato diferente", "Só aceito .xlsx ou .csv.")
                return False, f"{nome}: formato inválido"
            self.resultado = []
            n_loja = self._aplicar_loja()
            self._refresh()
            self.saude = dict(ULTIMOS_PEDIDOS) if ULTIMOS_PEDIDOS else {}
            self.arq_atual = arq.rsplit("/", 1)[-1].rsplit("\\", 1)[-1][:30]
            self.slot_ativo = None
            self._slot_estilos()
            if self.saude.get("total"):
                try:
                    tx = round(self.saude["cancelados"]/self.saude["total"]*100, 1)
                    self.p_dev.delete(0, "end"); self.p_dev.insert(0, str(tx).replace(".", ","))
                except Exception:
                    pass
            m = _re.search(r"(\d{8})_(\d{8})", arq)
            if m and (forcar_periodo or not self.p_de.get().strip()):
                try:
                    d1 = _dt.datetime.strptime(m.group(1), "%Y%m%d").strftime("%d/%m/%Y")
                    d2 = _dt.datetime.strptime(m.group(2), "%Y%m%d").strftime("%d/%m/%Y")
                    self.p_de.delete(0, "end"); self.p_de.insert(0, d1)
                    self.p_ate.delete(0, "end"); self.p_ate.insert(0, d2)
                except ValueError:
                    pass
            self.analisar()
            self.nb.select(self.tA)
            per = f"{self.p_de.get().strip()} a {self.p_ate.get().strip()}".strip(" a")
            resumo = f"{nome}: {len(self.dados)} produtos ({per or 'sem período'})"
            if n_loja:
                resumo += f" [{n_loja} dados restaurados]"
            if avisar:
                if self.saude.get("total"):
                    s = self.saude
                    messagebox.showinfo("Deu certo!",
                        f"Agreguei {s['validos']} vendas válidas em {len(self.dados)} produtos\n"
                        f"({s['cancelados']} cancelados + {s['nao_pagos']} não pagos separados).\n"
                        "Análise pronta na aba 2!")
                else:
                    messagebox.showinfo("Deu certo!", f"Li {len(self.dados)} produtos.\nAnálise pronta na aba 2!")
            return True, resumo
        except Exception as e:
            try:
                import traceback, os
                log = os.path.join(os.path.dirname(os.path.abspath(__file__)), "erro.log")
                with open(log, "a", encoding="utf-8") as f:
                    f.write(f"\n--- {arq} ---\n{traceback.format_exc()}")
            except Exception:
                pass
            if avisar:
                messagebox.showerror("Não consegui ler", str(e) + "\n\nGuardei o detalhe no arquivo erro.log da pasta do programa.")
            return False, f"{nome}: {e}"

    def exe(self):
        self.dados = exemplo(); self.resultado = []; self._aplicar_loja(); self._refresh()
        self.arq_atual = "exemplo"
        self.slot_ativo = None
        self._slot_estilos()
        self.analisar()
        self.nb.select(self.tA)

    def add(self):
        p = simpledialog.askstring("Produto", "Qual o nome do produto?")
        if not p: return
        q = simpledialog.askstring("Produto", "Quantas unidades vendeu?")
        v = simpledialog.askstring("Produto", "Faturou quanto no total? (ou o preço de um)")
        try:
            f = num(v); qtd = int(num(q))
            if 0 < f < qtd * 2 and qtd > 0:
                f = round(f*qtd, 2)
            self.dados.append({"produto": p, "qtd": qtd, "fat": f, "est": 0, "custo": 0})
            self.resultado = []; self._aplicar_loja(); self._refresh()
        except ValueError: messagebox.showerror("Ops", "Use só números, por favor.")

    def limpar(self):
        self.dados = []; self.resultado = []; self._refresh()

    def _periodo_txt(self):
        d1, d2 = self.p_de.get().strip(), self.p_ate.get().strip()
        if d1 or d2:
            return f"Período: {d1 or '?'} a {d2 or '?'}"
        return "Período: não informado (vale para os últimos dados exportados)"

    def analisar(self):
        if not self.dados:
            messagebox.showwarning("Calma lá", "Primeiro abra sua planilha no passo 1.")
            return
        self.resultado, self.total = classificar(self.dados)
        self._refresh(com_classe=True)
        self.out.delete("1.0", "end"); self.out.insert("end", self._texto_analise())
        self._status()
        self._hist_save()

    def _texto_analise(self):
        por = {"A": [x for x in self.resultado if x["classe"] == "A"],
               "B": [x for x in self.resultado if x["classe"] == "B"],
               "C": [x for x in self.resultado if x["classe"] == "C"]}
        n = len(self.resultado)
        fat = lambda c: sum(x["fat"] for x in por[c])
        top_n = max(1, round(n*0.2))
        conc = sum(x["fat"] for x in self.resultado[:top_n])/self.total*100 if self.total else 0
        L = ["SUA ANÁLISE — explicada de forma simples", self._periodo_txt(), "="*60,
             f"\nNo total você faturou {dinheiro(self.total)} com {n} produtos.",
             f"[A] Campeões: {len(por['A'])} produtos → {dinheiro(fat('A'))}",
             f"[B] Podem crescer: {len(por['B'])} produtos → {dinheiro(fat('B'))}",
             f"[C] Parados: {len(por['C'])} produtos → {dinheiro(fat('C'))}",
             f"\nPoucos produtos geram {conc:.0f}% do dinheiro. " +
             ("Cuide bem dos campeões: eles pagam as contas!" if conc >= 60 else "Dá para crescer em várias frentes.")]
        L.append("\nSEUS 5 CAMPEÕES (nunca deixe faltar):")
        for x in self.resultado[:5]:
            L.append(f"  • {x['produto'][:55]} — {dinheiro(x['fat'])} ({x['qtd']} vendidos)")
        L.append("\n[A] CAMPEÕES: estoque sempre cheio, fotos caprichadas, anuncie neles primeiro.")
        L.append("[B] DO MEIO: podem virar campeões. Melhore foto e título, teste um pouco de anúncio.")
        L.append("[C] PARADOS: venda junto com um campeão em kit com desconto. Parado há 2 meses? Liquide e não compre mais.")
        alertas = [f"URGENTE: {x['produto'][:50]} é CAMPEÃO com pouco estoque ({x['est']} un) — compre mais já!"
                   for x in por["A"] if 0 < x["est"] < x["qtd"]*0.3]
        alertas += [f"ATENÇÃO: {x['produto'][:50]} está parado com muito estoque ({x['est']} un) — monte um kit."
                    for x in por["C"] if x["est"] > max(x["qtd"]*2, 1)][:5]
        if alertas: L.append("\nAVISOS IMPORTANTES:\n" + "\n".join("  " + a for a in alertas))
        elif any(x["est"] for x in self.resultado): L.append("\nTudo certo com o estoque, sem problemas graves.")
        else: L.append("\nDica: se sua planilha tiver coluna de estoque, eu aviso quem vai faltar e quem encalhou.")
        return "\n".join(L)

    def copiar(self):
        self.clipboard_clear(); self.clipboard_append(self.out.get("1.0", "end")); self.update()
        messagebox.showinfo("Pronto", "Texto copiado!")

    def saverep(self):
        arq = filedialog.asksaveasfilename(defaultextension=".txt", filetypes=[("Texto", "*.txt")])
        if arq:
            with open(arq, "w", encoding="utf-8") as f: f.write(self.out.get("1.0", "end"))
            messagebox.showinfo("Pronto", f"Salvei em:\n{arq}")

    def savexlsx(self):
        if not self.resultado: messagebox.showwarning("Calma lá", "Faça a análise primeiro."); return
        arq = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")])
        if not arq: return
        try:
            import openpyxl
        except ImportError:
            messagebox.showerror("Ops", "Rode: pip install openpyxl"); return
        try:
            wb = openpyxl.Workbook()
            per = f"{self.p_de.get().strip()} a {self.p_ate.get().strip()}".strip(" a")
            ws = wb.active; ws.title = "Resumo"
            ws.append(["Curva", "Qtd produtos", "Faturamento", "% do total", "Período"])
            for cl in "ABC":
                g = [x for x in self.resultado if x.get("classe") == cl]
                f = sum(x["fat"] for x in g)
                ws.append([cl, len(g), round(f, 2), round(f/(self.total or 1)*100, 1), per])
            for cl, nome in [("A", "Curva A"), ("B", "Curva B"), ("C", "Curva C")]:
                wa = wb.create_sheet(nome)
                wa.append(["Produto", "Qtd vendida", "Faturamento", "% do total", "% acumulado", "Estoque", "Custo", "Período"])
                for x in [r for r in self.resultado if r.get("classe") == cl]:
                    wa.append([x["produto"], x["qtd"], x["fat"], x["part"], x["acum"],
                               x.get("est", 0), x.get("custo", 0), per])
                wa.append(["TOTAL", sum(x["qtd"] for x in self.resultado if x.get("classe") == cl),
                           round(sum(x["fat"] for x in self.resultado if x.get("classe") == cl), 2),
                           "", "", "", "", ""])
            for wa in wb.worksheets:
                for col in wa.columns:
                    wa.column_dimensions[col[0].column_letter].width = 22
                wa.column_dimensions["A"].width = 55
            wb.save(arq)
            messagebox.showinfo("Pronto", f"Salvei organizado por curvas em:\n{arq}")
        except Exception as e:
            messagebox.showerror("Ops", f"Não consegui salvar: {e}")

    def savecsv(self):
        if not self.resultado: messagebox.showwarning("Calma lá", "Faça a análise primeiro."); return
        arq = filedialog.asksaveasfilename(defaultextension=".csv", filetypes=[("CSV", "*.csv")])
        if not arq: return
        with open(arq, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=["produto", "qtd", "fat", "classe", "part", "acum", "periodo"], delimiter=";")
            w.writeheader()
            per = f"{self.p_de.get().strip()} a {self.p_ate.get().strip()}"
            for x in self.resultado:
                row = {k: x[k] for k in ("produto", "qtd", "fat", "classe", "part", "acum")}
                row["periodo"] = per
                w.writerow(row)
        messagebox.showinfo("Pronto", f"Salvei em:\n{arq}\n(Abre certinho no Excel.)")

    # ----- precificador -----
    def _tela_preco(self):
        t = self.tP
        self.dica(t, "Escolha a plataforma e diga quanto custa cada produto (fornecedor). Eu desconto as taxas reais de outubro/2026 "
                     "(comissão + fixo + frete) e mostro seu lucro de verdade e o preço ideal — os anúncios do passo 4 usam essa margem real.", "roas")
        f = ttk.Frame(t); f.pack(fill="x", padx=14, pady=6)
        f.columnconfigure(1, weight=1); f.columnconfigure(3, weight=1)
        ttk.Label(f, text="Precificar para:").grid(row=0, column=0, sticky="w", pady=4)
        self.p_plat = ttk.Combobox(f, values=["Shopee", "ML Clássico", "ML Premium"], width=12, state="readonly")
        self.p_plat.set("Shopee"); self.p_plat.grid(row=0, column=1, sticky="we", padx=6, pady=4)
        self.p_plat.bind("<<ComboboxSelected>>", lambda e: self.p_toggle())
        ttk.Label(f, text="Reserva p/ ads (%):").grid(row=0, column=2, sticky="w", pady=4)
        self.p_ads = ttk.Entry(f, width=10); self.p_ads.insert(0, "10"); self.p_ads.grid(row=0, column=3, sticky="we", padx=6, pady=4)
        ttk.Label(f, text="Lucro que quero (%):").grid(row=1, column=0, sticky="w", pady=4)
        self.p_lucro = ttk.Entry(f, width=10); self.p_lucro.insert(0, "30"); self.p_lucro.grid(row=1, column=1, sticky="we", padx=6, pady=4)
        ttk.Label(f, text="Comissão ML:").grid(row=1, column=2, sticky="w", pady=4)
        self.p_com_info = ttk.Label(f, text=""); self.p_com_info.grid(row=1, column=3, sticky="w", padx=6, pady=4)
        ttk.Label(f, text="Imposto (%):").grid(row=2, column=0, sticky="w", pady=4)
        self.p_imp = ttk.Entry(f, width=10); self.p_imp.insert(0, "4"); self.p_imp.grid(row=2, column=1, sticky="we", padx=6, pady=4)
        ttk.Label(f, text="Perdas/dev. (%):").grid(row=2, column=2, sticky="w", pady=4)
        self.p_dev = ttk.Entry(f, width=10); self.p_dev.insert(0, "3"); self.p_dev.grid(row=2, column=3, sticky="we", padx=6, pady=4)
        f2 = ttk.Frame(t); f2.pack(fill="x", padx=14, pady=2)
        f2.columnconfigure(1, weight=1); f2.columnconfigure(3, weight=1)
        self.p_cpf = tk.BooleanVar()
        self.ck_cpf = ttk.Checkbutton(f2, text="Sou CPF c/ +450 pedidos (Shopee +R$3)", variable=self.p_cpf)
        self.ck_cpf.grid(row=0, column=0, columnspan=2, sticky="w", pady=2)
        ttk.Label(f2, text="Peso do envio:").grid(row=0, column=2, sticky="w", pady=2)
        self.p_peso = ttk.Combobox(f2, values=["leve (até 0,5 kg)", "medio (0,5 a 2 kg)", "pesado (digito o frete)"], width=20, state="readonly")
        self.p_peso.set("leve (até 0,5 kg)"); self.p_peso.grid(row=0, column=3, sticky="we", padx=6, pady=2)
        self.p_flex = tk.BooleanVar()
        self.ck_flex = ttk.Checkbutton(f2, text="Envio Flex/por conta própria (ML)", variable=self.p_flex)
        self.ck_flex.grid(row=1, column=0, columnspan=2, sticky="w", pady=2)
        ttk.Label(f2, text="Frete manual R$:").grid(row=1, column=2, sticky="w", pady=2)
        self.p_frete = ttk.Entry(f2, width=10); self.p_frete.grid(row=1, column=3, sticky="we", padx=6, pady=2)
        g = ttk.Frame(t); g.pack(fill="x", padx=14, pady=4)
        ttk.Label(g, text="Custo do produto selecionado (R$):").pack(side="left")
        self.p_custo = ttk.Entry(g, width=12); self.p_custo.pack(side="left", padx=6)
        ttk.Button(g, text="Aplicar custo", command=self.p_aplicar).pack(side="left", padx=2)
        ttk.Button(g, text="Calcular tudo", style="Acao.TButton", command=self.p_calcular).pack(side="left", padx=8)
        ttk.Button(g, text="Ver tabela de taxas", command=self.taxas_popup).pack(side="left", padx=2)
        ttk.Button(g, text="Shopee x ML", command=self.comp_popup).pack(side="left", padx=2)
        self.ptree = self.tabela(t, [("produto", "PRODUTO", 300, True), ("preco", "PREÇO", 85, False),
                                     ("custo", "CUSTO", 85, False), ("taxas", "TAXAS", 85, False),
                                     ("frete", "FRETE", 85, False), ("margem", "LUCRO", 85, False),
                                     ("ideal", "PREÇO IDEAL", 100, False)], 8)
        self.ptree.tag_configure("A", background="#DFF5DF"); self.ptree.tag_configure("B", background="#FFF3D6")
        self.ptree.tag_configure("C", background="#FFE3E3")
        self.p_toggle()

    def p_comissao_ml(self):
        """Comissão automática pela tabela oficial: Clássico 12%, Premium 17% (médias por categoria)."""
        return 17.0 if "Premium" in self.p_plat.get() else 12.0

    def p_toggle(self):
        ml = self.p_plat.get() != "Shopee"
        self.p_com_info.config(text=f"{self.p_comissao_ml():.0f}% automática" if ml else "—")
        self.p_peso.config(state="readonly" if ml else "disabled")
        self.ck_flex.config(state="normal" if ml else "disabled")

    def _dias_periodo(self):
        try:
            d1 = _data_br(self.p_de.get()); d2 = _data_br(self.p_ate.get())
            if d1 and d2 and d2 >= d1:
                return max((d2-d1).days + 1, 1)
        except Exception:
            pass
        return 30

    def _set_estoque(self, prod, qtd):
        for base in (self.dados, self.resultado):
            for x in base:
                if x["produto"] == prod:
                    x["est"] = qtd

    def estoque_popup(self):
        if not self.dados:
            messagebox.showwarning("Calma lá", "Abra sua planilha primeiro (passo 1)."); return
        pop = tk.Toplevel(self)
        pop.title("Informar estoque")
        pop.geometry("480x260")
        ttk.Label(pop, text="Produto:").pack(anchor="w", padx=10, pady=(10, 0))
        cb = ttk.Combobox(pop, values=[x["produto"] for x in self.dados], width=55, state="readonly")
        cb.pack(fill="x", padx=10, pady=4)
        if self.dados:
            cb.set(self.dados[0]["produto"])
        ttk.Label(pop, text="Estoque atual (unidades):").pack(anchor="w", padx=10)
        ent = ttk.Entry(pop, width=14); ent.pack(anchor="w", padx=10, pady=4)

        def salvar(fechar=True):
            try:
                q = int(num(ent.get()))
            except ValueError:
                messagebox.showerror("Ops", "Use só números."); return
            prod = cb.get()
            self._set_estoque(prod, q)
            reg = self.loja.get(prod, {})
            reg["est"] = q
            self.loja[prod] = reg
            self._loja_save()
            self._refresh(bool(self.resultado))
            if fechar:
                messagebox.showinfo("Pronto", "Estoque salvo! Ele fica guardado para a próxima vez.")
                pop.destroy()
            else:
                nomes = [x["produto"] for x in self.dados]
                try:
                    cb.set(nomes[nomes.index(prod)+1])
                except (ValueError, IndexError):
                    pass
                ent.delete(0, "end")
                ent.focus()
        ttk.Button(pop, text="Salvar", style="Acao.TButton", command=salvar).pack(fill="x", padx=10, pady=(0, 4))
        ttk.Button(pop, text="Salvar e ir p/ próximo", command=lambda: salvar(False)).pack(fill="x", padx=10, pady=(0, 8))

    # ----- acompanhar ruptura -----
    def ruptura_popup(self):
        if not self.resultado:
            messagebox.showwarning("Calma lá", "Faça a análise (passo 2) primeiro."); return
        dias = self._dias_periodo()
        if not any(x.get("est", 0) for x in self.resultado):
            messagebox.showwarning("Sem estoque",
                "Você ainda não informou estoques.\nUse 'Informar estoque' e volte aqui."); return
        pop = tk.Toplevel(self)
        pop.title("Acompanhar ruptura de estoque")
        pop.geometry("760x600")
        tv = self.tabela(pop, [("produto", "PRODUTO", 320, True), ("cl", "CL", 45, False),
                               ("venda", "VENDA/DIA", 90, False), ("est", "ESTOQUE", 85, False),
                               ("cob", "COBERTURA", 100, False), ("rep", "REPÔR P/ 30D", 110, False)], 14)
        tv.tag_configure("CRIT", background="#FFCDD2")
        tv.tag_configure("AT", background="#FFF3D6")
        tv.tag_configure("OKK", background="#DFF5DF")
        crit = at = tot_rep = 0
        linhas = []
        for x in self.resultado:
            daily = x["qtd"]/dias if dias else 0
            est = x.get("est", 0) or 0
            if daily <= 0:
                continue
            cob = est/daily
            if cob < 7:
                st, tag = ("ZERADO" if est <= 0 else f"{cob:.0f} dias - CRÍTICO"), "CRIT"
                crit += 1
            elif cob < 14:
                st, tag = f"{cob:.0f} dias - ATENÇÃO", "AT"
                at += 1
            elif cob < 30:
                st, tag = f"{cob:.0f} dias", "OKK"
            else:
                st, tag = f"{cob:.0f} dias - ok", ""
            rep = max(0, int(30*daily - est + 0.5))
            tot_rep += rep
            linhas.append((cob, x["produto"][:55], x.get("classe", ""), f"{daily:.1f}", est, st, rep, tag))
        linhas.sort()
        for _cob, prod, cl, daily, est, st, rep, tag in linhas:
            tv.insert("", "end", values=(prod, cl, daily, est, st, rep if rep else "—"), tags=(tag,) if tag else ())
        tk.Label(pop, text=f"CRÍTICOS (<7 dias): {crit}   |   ATENÇÃO (<14 dias): {at}   |   Total a repor p/ 30 dias: {tot_rep} un",
                 bg=FUNDO, font=("Segoe UI", 11, "bold")).pack(fill="x", padx=10, pady=(0, 8))

    def _set_custo(self, prod, custo):
        for base in (self.dados, self.resultado):
            for x in base:
                if x["produto"] == prod:
                    x["custo"] = custo

    def taxas_popup(self):
        pop = tk.Toplevel(self)
        pop.title("Tabela de taxas Shopee e ML (out/2026)")
        pop.geometry("660x650")
        txt = scrolledtext.ScrolledText(pop, wrap="word", height=9, font=("Segoe UI", 11))
        txt.pack(fill="both", expand=True, padx=10, pady=(10, 4))
        txt.insert("end",
            "SHOPEE — comissão + taxa fixa por item (desde 01/10/2026)\n"
            "• Até R$ 79,99: 20% + R$ 4,50 por item\n"
            "• R$ 80 a 99,99: 14% + R$ 16,00 por item\n"
            "• R$ 100 a 199,99: 14% + R$ 20,00 por item\n"
            "• R$ 200 ou mais: 14% + R$ 26,00 por item\n"
            "• Abaixo de R$ 9: taxa = metade do preço\n"
            "• CPF com +450 pedidos em 90 dias: +R$ 3,00 por item\n"
            "• Frete Grátis já embutido (subsídios de R$ 20/30/40 por faixa)\n"
            "• Comissão só sobre o produto; cancelado/devolvido não paga\n"
            "\nMERCADO LIVRE — comissão + envio (desde 02/03/2026)\n"
            "• Clássico 10-14% (uso 12%) · Premium 15-19% (uso 17%)\n"
            "• Abaixo de R$ 79 no Flex/próprio: fixo R$ 7,75 por item\n"
            "• Full/Coleta/Agência: tabela por peso e faixa (ex: 0,5kg = R$ 8,15 até R$ 79; R$ 12,95 de R$ 79-100)\n"
            "• Acima de R$ 79: frete grátis obrigatório, pago por você\n")
        txt.config(state="disabled")
        tk.Label(pop, text="1 — QUANTO RECEBO DE UMA VENDA", bg=FUNDO, fg=ROSA_ESCURO,
                 font=("Segoe UI", 12, "bold"), anchor="w").pack(fill="x", padx=10, pady=(6, 0))
        f = ttk.Frame(pop); f.pack(fill="x", padx=10, pady=4)
        ttk.Label(f, text="Simular preço R$:").pack(side="left")
        ent = ttk.Entry(f, width=12); ent.pack(side="left", padx=6)
        combo = ttk.Combobox(f, values=["Shopee", "ML Clássico", "ML Premium"], width=12, state="readonly")
        combo.set("Shopee"); combo.pack(side="left", padx=2)
        out = tk.Label(pop, text="O resultado aparece aqui.", bg=ROSA_CLARO, fg=CINZA_TXT,
                         font=("Segoe UI", 12, "bold"), justify="left")
        out.pack(fill="x", padx=10, pady=(0, 4))

        def simular():
            try: p = num(ent.get())
            except ValueError: p = 0
            if p <= 0:
                out.config(text="Digite um preço válido."); return
            c = combo.get()
            if c == "Shopee":
                if p < 9: pct_v, fixa = p*0.5, 0.0
                elif p <= 79.99: pct_v, fixa = p*0.20, 4.50
                elif p < 100: pct_v, fixa = p*0.14, 16.0
                elif p < 200: pct_v, fixa = p*0.14, 20.0
                else: pct_v, fixa = p*0.14, 26.0
                out.config(text=f"Shopee {dinheiro(p)}:\nComissão % = {dinheiro(pct_v)}\nTaxa fixa = {dinheiro(fixa)}\n"
                               f"TOTAL taxas = {dinheiro(pct_v+fixa)}\nVocê recebe = {dinheiro(p-pct_v-fixa)}")
            else:
                com = 0.17 if "Premium" in c else 0.12
                cv, fr = p*com, envio_ml(p, "leve")
                out.config(text=f"{c} {dinheiro(p)} (leve, Full):\nComissão {com*100:.0f}% = {dinheiro(cv)}\n"
                               f"Envio = {dinheiro(fr)}\nTOTAL taxas = {dinheiro(cv+fr)}\nVocê recebe = {dinheiro(p-cv-fr)}")
        ttk.Button(f, text="Calcular", style="Acao.TButton", command=simular).pack(side="left", padx=6)
        tk.Label(pop, text="2 — POR QUANTO ANUNCIAR (custo + lucro)", bg=FUNDO, fg=ROSA_ESCURO,
                 font=("Segoe UI", 12, "bold"), anchor="w").pack(fill="x", padx=10, pady=(2, 0))
        f2 = ttk.Frame(pop); f2.pack(fill="x", padx=10, pady=4)
        ttk.Label(f2, text="Ou descubra o preço: meu custo R$:").pack(side="left")
        ent_c = ttk.Entry(f2, width=10); ent_c.pack(side="left", padx=4)
        ttk.Label(f2, text="Quero ganhar %:").pack(side="left")
        ent_l = ttk.Entry(f2, width=8); ent_l.insert(0, "30"); ent_l.pack(side="left", padx=4)
        out2 = tk.Label(pop, text="O preço ideal aparece aqui.", bg=ROSA_CLARO, fg=CINZA_TXT,
                          font=("Segoe UI", 12, "bold"), justify="left")
        out2.pack(fill="x", padx=10, pady=(0, 10))

        def ideal():
            try:
                cu, lu = num(ent_c.get()), num(ent_l.get())/100
            except ValueError:
                cu, lu = 0, 0
            if cu <= 0 or lu <= 0:
                out2.config(text="Digite custo e % de ganho."); return
            c = combo.get()
            ta = 0.10
            if c == "Shopee":
                p = preco_ideal_shopee(cu, ta, lu)
                tx = taxa_shopee(p)
            else:
                com = 0.17 if "Premium" in c else 0.12
                p = preco_ideal_ml(cu, ta, lu, com)
                tx = taxa_ml(p, "Premium" in c)
            if not p:
                out2.config(text="Conta não fecha: diminua o ganho ou a reserva."); return
            sobra = p - cu - tx
            out2.config(text=f"Anuncie por {dinheiro(p)}:\nTaxas = {dinheiro(tx)} | Sobra {dinheiro(sobra)} "
                             f"({sobra/p*100:.0f}%: {lu*100:.0f}% lucro + 10% ads)")
        ttk.Button(f2, text="Preço ideal", style="Acao.TButton", command=ideal).pack(side="left", padx=6)

    def p_aplicar(self):
        sel = self.ptree.selection()
        if not sel:
            messagebox.showwarning("Selecione", "Clique num produto da lista primeiro."); return
        try: custo = num(self.p_custo.get())
        except ValueError:
            messagebox.showerror("Ops", "Use só números no custo."); return
        prod = self.ptree.item(sel[0])["values"][0]
        achou = False
        for x in self.dados + self.resultado:
            if x["produto"] == prod:
                x["custo"] = custo; achou = True
        if not achou:
            messagebox.showwarning("Ops", "Não achei esse produto. Tente de novo."); return
        reg = self.loja.get(prod, {})
        reg["custo"] = custo
        self.loja[prod] = reg
        self._loja_save()
        self.p_calcular()

    def p_detalhe(self, preco):
        """(taxas totais, parte do frete) pelas regras de cada plataforma."""
        plat = self.p_plat.get()
        try: frete_man = num(self.p_frete.get())
        except ValueError: frete_man = 0.0
        if plat == "Shopee":
            # Frete Grátis embutido na comissão; vendedor só paga extra manual.
            return taxa_shopee(preco, cpf_extra=self.p_cpf.get(), frete_extra=frete_man), frete_man
        com = self.p_comissao_ml()/100
        peso_nome = self.p_peso.get()
        peso = {"leve (até 0,5 kg)": "leve", "medio (0,5 a 2 kg)": "medio"}.get(peso_nome, "leve")
        if peso_nome.startswith("pesado"):
            frete = frete_man  # pesado: vale o que você digitar
        elif self.p_flex.get() and preco < 79:
            frete = 7.75  # Flex/próprio abaixo de R$79: custo fixo
        else:
            frete = envio_ml(preco, peso)  # Full/Coleta/Agência ou frete grátis acima de R$79
        return preco*com + frete, frete

    def p_taxas(self, preco):
        """Taxas em R$ conforme plataforma escolhida no precificador."""
        total, _frete = self.p_detalhe(preco)
        return total

    def p_imposto(self):
        try: imp = num(self.p_imp.get())/100
        except (ValueError, AttributeError): imp = 0.04
        try: dev = num(self.p_dev.get())/100
        except (ValueError, AttributeError): dev = 0.03
        return imp, dev

    def p_ideal(self, custo, ta, mlucro, plat):
        """Preço ideal usando as funções oficiais por plataforma (com imposto e perdas)."""
        try:
            imp, dev = self.p_imposto()
            if plat == "Shopee":
                return preco_ideal_shopee(custo, ta, mlucro, self.p_cpf.get(), num(self.p_frete.get() or 0), imp, dev)
            peso_nome = self.p_peso.get()
            peso = {"leve (até 0,5 kg)": "leve", "medio (0,5 a 2 kg)": "medio"}.get(peso_nome, "leve")
            if peso_nome.startswith("pesado"):
                peso = "manual"
            return preco_ideal_ml(custo, ta, mlucro, self.p_comissao_ml()/100, peso,
                                  self.p_flex.get(), num(self.p_frete.get() or 0), imp, dev)
        except (ValueError, TypeError):
            return 0.0

    def p_calcular(self):
        if not self.dados:
            messagebox.showwarning("Calma lá", "Primeiro abra sua planilha no passo 1."); return
        try:
            ta = num(self.p_ads.get())/100; ml = num(self.p_lucro.get())/100
            imp, dev = self.p_imposto()
        except ValueError:
            messagebox.showerror("Ops", "Use só números nas taxas."); return
        plat = self.p_plat.get()
        for i in self.ptree.get_children(): self.ptree.delete(i)
        for x in sorted(self.dados, key=lambda d: d["fat"], reverse=True):
            pm = x["fat"]/x["qtd"] if x["qtd"] else 0
            cu = x.get("custo", 0) or 0
            tx, fr = self.p_detalhe(pm) if pm else (0, 0)
            margem = ((pm - cu - tx)/pm - imp - dev if pm else 0)
            ideal = self.p_ideal(cu, ta, ml, plat)
            cor = "A" if cu and margem >= ml else ("B" if cu and margem >= ml/2 else "C")
            self.ptree.insert("", "end", values=(x["produto"], dinheiro(pm), dinheiro(cu), dinheiro(tx), dinheiro(fr),
                f"{margem*100:.0f}%" if cu else "sem custo", dinheiro(ideal) if ideal else "—"), tags=(cor,))
        self.p_cfg = {"plat": plat, "cpf": self.p_cpf.get(), "flex": self.p_flex.get(),
                      "peso": self.p_peso.get(), "com": self.p_comissao_ml(), "frete": self.p_frete.get()}
        messagebox.showinfo("Pronto", f"Preços calculados para {plat}!\nVerde = lucro bom. O passo 4 agora usa o lucro real de cada produto.")

    # ----- verba -----
    def _roas_real(self, x, c, plat):
        """M0 = margem antes dos ads. Meta = 1/(M0 - lucro líquido). Retorna (M0, alvo, preco, real, be, inviavel)."""
        preco = x["fat"]/x["qtd"] if x["qtd"] else 0
        custo = x.get("custo", 0) or 0
        try:
            L = num(self.v_margem.get())/100
        except (ValueError, AttributeError):
            L = 0.22
        if preco > 0 and 0 < custo < preco:
            try:
                imp, dev = self.p_imposto()
                cfg = getattr(self, "p_cfg", None) or {}
                if plat == "Shopee":
                    fee = taxa_shopee(preco, cpf_extra=cfg.get("cpf", False),
                                      frete_extra=num(cfg.get("frete", 0) or 0))
                else:
                    is_prem = ("Premium" in (cfg.get("plat", "") or "")) or \
                              (not cfg.get("plat") and "Premium" in self.v_mltipo.get())
                    scfg = cfg.get("peso", "")
                    peso = {"leve (até 0,5 kg)": "leve", "medio (0,5 a 2 kg)": "medio"}.get(scfg, "leve")
                    com_df = cfg.get("com", 0) or (17.0 if is_prem else 12.0)
                    if scfg.startswith("pesado"):
                        fee = preco*(num(com_df)/100) + num(cfg.get("frete", 0) or 0)
                    else:
                        fee = taxa_ml(preco, premium=is_prem, com_pct=num(com_df),
                                      peso=peso, flex=cfg.get("flex", False))
                M0 = (preco - custo - fee)/preco - imp - dev
                if M0 <= 0.02:
                    return M0, 0, preco, True, 0, True
                room = M0 - L
                if room <= 0.02:
                    return M0, 0, preco, True, round(1/M0, 1), True
                return M0, round(1/room, 1), preco, True, round(1/M0, 1), False
            except (ValueError, TypeError, AttributeError):
                pass
        return 0, c["roas"], preco, False, 0, False

    def v_gerar(self):
        if not self.dados:
            messagebox.showwarning("Calma lá", "Primeiro abra sua planilha no passo 1.")
            return
        if not self.resultado:
            self.resultado, self.total = classificar(self.dados); self._refresh(com_classe=True)
        _c = {d["produto"]: d.get("custo", 0) for d in self.dados}
        for _x in self.resultado:
            _x["custo"] = _c.get(_x["produto"], _x.get("custo", 0))
        try:
            vt = num(self.v_total.get()); mg = num(self.v_margem.get())/100
            pa, pb, pc = num(self.v_pa.get()), num(self.v_pb.get()), num(self.v_pc.get())
        except ValueError:
            messagebox.showerror("Ops", "Use só números nos campos de valores."); return
        if vt <= 0:
            messagebox.showwarning("Faltou a verba", "Diga quanto pode gastar por dia (exemplo: 50)."); return
        if mg <= 0:
            messagebox.showwarning("Faltou o lucro", "Diga seu lucro em % (exemplo: 30)."); return
        if pa + pb + pc != 100:
            messagebox.showwarning("Soma errada", f"A + B + C está dando {pa+pb+pc:.0f}%. Ajuste para somar 100."); return
        sel = self.v_plat.get()
        plats = ["Shopee", "Mercado Livre"] if sel == "Nos dois" else [sel]
        shares = {"Shopee": 0.5, "Mercado Livre": 0.5} if len(plats) == 2 else {plats[0]: 1.0}
        imp, dev = self.p_imposto()
        try: luc = num(self.v_margem.get())/100
        except (ValueError, AttributeError): luc = 0.22
        if luc <= 0:
            messagebox.showwarning("Faltou o lucro", "Diga seu lucro líquido final em % (exemplo: 22)."); return
        cfg = {}
        for p in plats:
            com = 17.0 if (p == "Mercado Livre" and "Premium" in self.v_mltipo.get()) else COMISSAO_PADRAO[p]
            cpc = CPC_PADRAO[p]
            # Sem custo informado: regra de especialista = meta 1 ÷ lucro desejado
            cfg[p] = {"com": com, "cpc": cpc, "mgl": 0, "be": 0, "roas": round(1/luc, 1),
                      "conv": CONV_PADRAO[p], "verba": vt*shares[p]}
        por = {"A": [x for x in self.resultado if x["classe"] == "A"],
               "B": [x for x in self.resultado if x["classe"] == "B"],
               "C": [x for x in self.resultado if x["classe"] == "C"]}
        fa = sum(x["fat"] for x in por["A"]) or 1
        fb = sum(x["fat"] for x in por["B"]) or 1
        for i in self.vtree.get_children(): self.vtree.delete(i)
        self.vtree.tag_configure("CAMP", background=AZUL, foreground="white")
        self.vtree.tag_configure("ALERT", background="#FFE3E3")
        L = []
        ret_total = 0.0
        inviaveis = []
        nA = por["A"][:8]
        como = {"Shopee": "Anúncio de Busca (aparece quando pesquisam) + Descoberta (aparece para quem olha parecidos).",
                "Mercado Livre": "Product Ads (seu produto em destaque na busca). Precisa de estoque e envio ativo."}
        for p in plats:
            c = cfg[p]
            vA, vB = c["verba"]*pa/100, c["verba"]*pb/100
            L.append("="*64)
            L.append(f"PLANO DE MÍDIA — {p.upper()}  |  verba {dinheiro(c['verba'])}/dia  |  meta {c['roas']:.1f}x")
            L.append(f"Como anunciar: {como[p]}")
            if c["roas"] > 10:
                L.append("ALERTA: meta acima de 10x é quase inviável — baixe o lucro, suba preços ou venda em kit.")
            L.append("-"*64)
            L.append(f"CAMPANHA 1 — CAMPEÕES · 1 anúncio por produto (unitário) · {dinheiro(vA)}/dia")
            self.vtree.insert("", "end", values=(f"CAMPANHA 1 · {p[:2]} · Campeões · {dinheiro(vA)}/dia", "", "", "", "", ""), tags=("CAMP",))
            for x in nA:
                vb = vA*x["fat"]/fa
                M0, alvo, preco, real, be, inv = self._roas_real(x, c, p)
                if inv:
                    inviaveis.append(x["produto"][:50])
                    self.vtree.insert("", "end", values=(x["produto"][:55], p[:2], "Sozinho", dinheiro(0), "inviável", dinheiro(0)), tags=("ALERT",))
                    L.append(f"  × {x['produto'][:50]} — SEM VERBA (não comporta {luc*100:.0f}% líquido)")
                    continue
                ret = vb*alvo; ret_total += ret
                rotulo = f"{alvo:.1f}x (eq {be:.1f}x)" if real else f"{alvo:.1f}x"
                self.vtree.insert("", "end", values=(x["produto"][:55], p[:2], "Sozinho", dinheiro(vb), rotulo, dinheiro(ret)))
                L.append(f"  • {x['produto'][:50]} — {dinheiro(vb)}/dia → meta {rotulo} → retorno {dinheiro(ret)}/dia")
                if real and preco*M0*c["conv"] < c["cpc"]:
                    L.append(f"    Cuidado: clique ~{dinheiro(c['cpc'])} acima do teto {dinheiro(preco*M0*c['conv'])}.")
            L.append("-"*64)
            L.append(f"CAMPANHA 2 — TESTES · produtos B, 1 anúncio por produto · {dinheiro(vB)}/dia")
            self.vtree.insert("", "end", values=(f"CAMPANHA 2 · {p[:2]} · Testes ({len(por['B'])} produtos) · {dinheiro(vB)}/dia", "", "", "", "", ""), tags=("CAMP",))
            for x in por["B"]:
                vb = vB*x["fat"]/fb
                M0, alvo, preco, real, be, inv = self._roas_real(x, c, p)
                if inv:
                    inviaveis.append(x["produto"][:50])
                    self.vtree.insert("", "end", values=(x["produto"][:55], p[:2], "Sozinho", dinheiro(0), "inviável", dinheiro(0)), tags=("ALERT",))
                    L.append(f"  × {x['produto'][:50]} — SEM VERBA (não comporta {luc*100:.0f}% líquido)")
                    continue
                r2 = round(alvo*0.9, 1); ret = vb*r2; ret_total += ret
                rotulo = f"{r2:.1f}x (eq {be:.1f}x)" if real else f"{r2:.1f}x"
                self.vtree.insert("", "end", values=(x["produto"][:55], p[:2], "Sozinho", dinheiro(vb), rotulo, dinheiro(ret)))
                L.append(f"  • {x['produto'][:50]} — {dinheiro(vb)}/dia → meta {rotulo} → retorno {dinheiro(ret)}/dia")
            if (vA and vA/max(len(nA), 1) < 5) or (vB and por["B"] and vB/len(por["B"]) < 5):
                L.append("  ATENÇÃO: verba diluída (< R$ 5/dia por anúncio). Junte produtos ou aumente a verba.")
            L.append("-"*64)
            nK = [x for x in por["C"] if x["fat"] > 0][:5]
            L.append(f"CAMPANHA 3 — KITS · vários produtos por anúncio · {len(nK)} kits · sem verba própria")
            self.vtree.insert("", "end", values=(f"CAMPANHA 3 · {p[:2]} · Kits · sem verba própria", "", "", "", "", ""), tags=("CAMP",))
            L.append("-"*64)
            L.append(f"PLANO DOS CURVA C — como melhorar a performance no {p} (sem verba solo, crescem via kit)")
            topsC = [t["produto"] for t in por["A"][:3]]
            for j, x in enumerate([c for c in por["C"] if c["fat"] > 0][:5]):
                M0, _al, _pr, real, _be, _iv = self._roas_real(x, c, p)
                acoes = []
                if len(x["produto"]) < 40:
                    acoes.append("esticar o título para 80+ letras")
                if real and M0 < 0.15:
                    acoes.append(f"margem apertada ({M0*100:.0f}%): subir preço ou trocar fornecedor")
                if not real:
                    acoes.append("informar o custo no passo 3")
                if x.get("est", 0) > max(x["qtd"]*2, 1):
                    acoes.append(f"estoque alto ({x['est']} un): girar em kit")
                par = topsC[j % len(topsC)][:35] if topsC else "campeão"
                acoes.append(f"anunciar SÓ no kit com {par}")
                L.append(f"  • {x['produto'][:50]} → {'; '.join(acoes)}.")
        for x in por["C"]:
            self.vtree.insert("", "end", values=(x["produto"][:55], "SH+ML", "Só em kit", dinheiro(0), f"{min(c['roas'] for c in cfg.values())*0.8:.1f}x", dinheiro(0)))
        L.append("="*64)
        L.append(f"RESUMO: gastando {dinheiro(vt)}/dia, a meta é trazer {dinheiro(ret_total)}/dia em vendas.")
        if inviaveis:
            L.append(f"SEM VERBA ({len(set(inviaveis))}): não comportam {luc*100:.0f}% líquido — suba o preço ou aceite lucro menor neles.")
        L.append("")
        L.append("ROTINA SEMANAL DO ESPECIALISTA:")
        L.append("1. Segunda: pausar anúncio com 30+ cliques e 0 vendas.")
        L.append("2. Quarta: dobrar a verba do que vendeu 3 dias seguidos.")
        L.append("3. Sexta: trocar foto/título de quem está com pouco clique.")
        L.append("4. Nunca mexer em público ou lance antes de 7 dias.")
        L.append("")
        L.append("KITS SUGERIDOS (parado + campeão, 10 a 15% off):")
        tops = [x["produto"][:40] for x in por["A"][:3]]
        for i, x in enumerate([c for c in por["C"] if c["fat"] > 0][:5], 1):
            L.append(f"  Kit {i}: {tops[i % len(tops)] if tops else 'seu campeão'} + {x['produto'][:40]}")
        self.v_out.delete("1.0", "end"); self.v_out.insert("end", "\n".join(L))
        self._status()

if __name__ == "__main__":
    App().mainloop()
